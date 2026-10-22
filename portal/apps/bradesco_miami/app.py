"""Transactions Offshore (Streamlit): Bradesco Miami, BTG US e Morgan Stanley (MS)

Tratamento dos arquivos Transactions_<banco> (mesmas regras para os dois bancos; só muda a pasta, o nome do
arquivo e o que a coluna Client mostra quando a conta não está no De-Para):
  1. apaga as linhas iniciais (1 a 6), e o cabeçalho passa a ser a linha 1;
  2. localiza "Date" e "Settlement Date" pelo nome (não pela posição) e as move para A e B;
  3. remove colunas vazias, sem deixar lacunas;
  4. apaga a coluna "Account Nickname";
  5. mantém só as linhas cuja "Activity Description" contém Buy ou Sell;
  6. mantém os Settlement Date de D-2 (dias úteis, sem feriados do Brasil) a partir da data do arquivo em diante, inclusive datas futuras; Settlement Date vazio também entra;
  7. descarta todas as colunas depois de "Details" (ela é mantida);
  8. cria a coluna "Client" (primeira coluna) buscando a Account na aba "De-Para Contas" do modelo
     (a mesma busca da planilha do modelo: Bradesco Miami escreve "erro" se a conta não existir; BTG US deixa vazio).

Morgan Stanley (MS e MS2) segue as mesmas regras, ajustadas ao layout do arquivo: Transaction Date e Activity Date
fazem o papel de Date e Settlement Date, o filtro é Sold/Bought na coluna "Activity", não há coluna Details (todas
ficam) e o Client busca os últimos 10 caracteres da Account no "Account Code" do De-Para (conta não achada: vazio).

Colunas de data saem como data abreviada (no Excel, o formato de data curta do sistema). Na aba Todos, o botão
"Montar e-mail no Outlook" abre o rascunho do Daily Report com uma tabela por banco (só os que têm linhas).
"""
import html
import io
import os
import re
import sys
import time
from datetime import date, datetime
from pathlib import Path
from urllib.parse import unquote

import pandas as pd
from dateutil.easter import easter
import streamlit as st
from openpyxl import load_workbook

# O portal roda este arquivo via st.navigation: garante a raiz (tema) no sys.path.
_RAIZ = Path(__file__).resolve().parents[2]
if str(_RAIZ) not in sys.path:
    sys.path.insert(0, str(_RAIZ))
from tema import aplicar_tema, topo  # noqa: E402

RAIZ_ARQUIVOS = Path(
    r"G:\Drives compartilhados\BO_OFF\Rotinas\Rotinas Diárias"
    r"\Segunda Rotina - Daily Report Offshore - Funds Settlements\Histórico Excel"
)
ARQUIVO_MODELO = RAIZ_ARQUIVOS.parent / "Modelo BTG_US BTG_Cayman MS.xlsm"
ABA_DEPARA = "De-Para Contas"
COLUNA_DEPARA_CONTA, COLUNA_DEPARA_CLIENTE = "Number", "Account"  # colunas B e E da aba
# Por banco: pasta dentro de <ano>/<Mês>, início do nome do arquivo e o Client quando a conta não está no De-Para
FILTRO_MS = r"(?<![A-Za-z])(?:sold|bought)(?![A-Za-z])"  # palavra inteira: Sold ou Bought
BANCOS = {
    "Bradesco Miami": {"pasta": "Bradesco Miami", "prefixo": "Transactions_Bradesco_Miami", "sem_cliente": "erro"},
    "BTG US": {"pasta": "BTG US", "prefixo": "Transactions_BTG_US", "sem_cliente": ""},
    # As duas exportações do Morgan Stanley ficam na mesma pasta ("Transactions_MS - " e "Transactions_MS2 - ") e são
    # tratadas e exibidas juntas
    "MS": {"pasta": "MS", "prefixos": ["Transactions_MS", "Transactions_MS2"], "sem_cliente": "", "tipo": "ms",
           # layout do arquivo do MS: Transaction Date e Activity Date fazem o papel de Date e Settlement Date;
           # o filtro usa Sold/Bought na coluna Activity; não há coluna Details (todas as colunas ficam)
           "iniciais": ["Transaction Date", "Activity Date"], "filtro": ("Activity", FILTRO_MS), "ultima": None},
}
FORMATO_DATA_EXCEL = "mm-dd-yy"  # formato interno do Excel para "data abreviada": segue a configuração regional
CHARS_CONTA_MS = 10  # RIGHT(Account, 10) da fórmula do modelo
COLUNA_DEPARA_CODIGO = "Account Code"  # coluna D do De-Para, usada na busca do MS
MESES = ["Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho", "Julho", "Agosto", "Setembro", "Outubro",
         "Novembro", "Dezembro"]
LINHAS_INICIAIS = 6
COLUNAS_INICIAIS = ["Date", "Settlement Date"]
COLUNA_APAGAR = "Account Nickname"
COLUNA_FILTRO = "Activity Description"
ULTIMA_COLUNA = "Details"  # o que vier depois dela é ignorado
DIAS_UTEIS_ATRAS = 2  # janela do Settlement Date: da data mais recente até D-2


def feriados_brasil(anos: range) -> list[pd.Timestamp]:
    """Feriados nacionais do Brasil (os mesmos do calendário ANBIMA): fixos e os que dependem da Páscoa."""
    datas = []
    for ano in anos:
        pascoa = pd.Timestamp(easter(ano))
        datas += [pd.Timestamp(ano, m, d) for m, d in
                  [(1, 1), (4, 21), (5, 1), (9, 7), (10, 12), (11, 2), (11, 15), (12, 25)]]
        if ano >= 2024:
            datas.append(pd.Timestamp(ano, 11, 20))  # Consciência Negra (nacional a partir de 2024)
        datas += [pascoa - pd.Timedelta(days=48), pascoa - pd.Timedelta(days=47),  # Carnaval (segunda e terça)
                  pascoa - pd.Timedelta(days=2),  # Sexta-feira Santa
                  pascoa + pd.Timedelta(days=60)]  # Corpus Christi
    return datas


DIA_UTIL = pd.offsets.CustomBusinessDay(holidays=feriados_brasil(range(2020, 2036)))  # pula fds e feriados
FILTRO_TEXTO = r"(?<![A-Za-z])(?:buy|sell)(?![A-Za-z])"  # palavra inteira, sem diferenciar maiúsculas de minúsculas

# set_page_config fica no portal (home.py)
aplicar_tema()


def _norm(nome) -> str:
    return " ".join(str(nome).lower().split())


def _conta_do_topo(topo_arq: pd.DataFrame) -> str | None:
    """Conta do cabeçalho do arquivo ("Account: 3GM035855"); None se for ALL_ACCOUNTS ou não existir."""
    for v in topo_arq.to_numpy().ravel():
        achou = re.match(r"\s*Account:\s*(\S+)", str(v))
        if achou and achou.group(1).upper() != "ALL_ACCOUNTS":
            return achou.group(1)
    return None


def _achar_coluna(colunas, nome: str):
    """Acha a coluna pelo nome, ignorando maiúsculas/minúsculas e espaços extras."""
    alvo = _norm(nome)
    for c in colunas:
        if _norm(c) == alvo:
            return c
    raise ValueError(f"Coluna '{nome}' não encontrada. Colunas: {[str(c) for c in colunas]}")


def _chave_conta(valor) -> str:
    """Número da conta como texto (no De-Para há contas em texto e em número)."""
    return str(valor).strip().casefold()


@st.cache_data(ttl=300, show_spinner=False)
def carregar_depara() -> dict[str, dict[str, str]]:
    """Lê a aba De-Para Contas do modelo e devolve dois mapas conta -> cliente:
    "numero" (coluna Number: Bradesco e BTG) e "codigo" (coluna Account Code: Morgan Stanley).
    Em chave repetida vale a primeira linha, como no XLOOKUP/VLOOKUP."""
    wb = load_workbook(ARQUIVO_MODELO, read_only=True, data_only=True)
    try:
        linhas = wb[ABA_DEPARA].iter_rows(values_only=True)
        cab = [_norm(c) for c in next(linhas)]
        i_cli = cab.index(_norm(COLUNA_DEPARA_CLIENTE))
        indices = {"numero": cab.index(_norm(COLUNA_DEPARA_CONTA)), "codigo": cab.index(_norm(COLUNA_DEPARA_CODIGO))}
        mapas = {nome: {} for nome in indices}
        for r in linhas:
            cliente = "" if r[i_cli] is None else str(r[i_cli])
            for nome, i in indices.items():
                if r[i] is not None and str(r[i]).strip():
                    mapas[nome].setdefault(_chave_conta(r[i]), cliente)
    finally:
        wb.close()
    return mapas


def tratar(arquivo, depara: dict[str, dict[str, str]], cfg: dict, mais_recente=None):
    sem_cliente = cfg["sem_cliente"]
    ms = cfg.get("tipo") == "ms"
    iniciais_cfg = cfg.get("iniciais", COLUNAS_INICIAIS)
    col_filtro, regex_filtro = cfg.get("filtro", (COLUNA_FILTRO, FILTRO_TEXTO))
    ultima_cfg = cfg.get("ultima", ULTIMA_COLUNA)
    # openpyxl lê a grade exatamente como está no Excel (o pandas descarta linhas vazias do topo)
    ws = load_workbook(arquivo, read_only=True, data_only=True).worksheets[0]
    grade = pd.DataFrame(list(ws.iter_rows(values_only=True)))
    conta_topo = _conta_do_topo(grade.iloc[:LINHAS_INICIAIS])
    bruto = grade.iloc[LINHAS_INICIAIS:].reset_index(drop=True)
    bruto.columns = bruto.iloc[0]  # a primeira linha restante é o cabeçalho
    df = bruto.iloc[1:].reset_index(drop=True)
    df.columns.name = None

    df = df.loc[:, df.columns.notna()]  # colunas sem título (ex.: coluna A vazia)
    df = df.dropna(axis=0, how="all").reset_index(drop=True)
    if not df.empty:  # arquivo só com cabeçalho (dia sem movimento): mantém as colunas
        df = df.dropna(axis=1, how="all")

    iniciais = [_achar_coluna(df.columns, n) for n in iniciais_cfg]
    for c in iniciais:  # alguns arquivos trazem texto vazio ("") no lugar de célula vazia
        df[c] = pd.to_datetime(df[c].replace(r"^\s*$", None, regex=True), errors="coerce")
    df = df[iniciais + [c for c in df.columns if c not in iniciais]]

    df = df.drop(columns=[c for c in df.columns if _norm(c) == _norm(COLUNA_APAGAR)])

    # Arquivo de uma conta só não traz a coluna Account: usa a conta do cabeçalho ("Account: 3GM...")
    if not any(_norm(c) == "account" for c in df.columns) and conta_topo:
        df.insert(len(iniciais_cfg), "Account", conta_topo)

    # Janela do Settlement Date, a partir da data mais recente da planilha (antes do filtro Buy/Sell)
    liquidacao = pd.to_datetime(df[iniciais[1]], errors="coerce")
    if mais_recente is None:  # sem referência externa: a data mais recente deste arquivo
        mais_recente = liquidacao.max().normalize() if liquidacao.notna().any() else None
    if mais_recente is not None and pd.notna(mais_recente):
        limite = mais_recente - DIAS_UTEIS_ATRAS * DIA_UTIL
        dia = liquidacao.dt.normalize()
        # Settlement Date vazio e liquidações futuras (depois da data do arquivo) também entram
        df = df[liquidacao.isna() | (dia >= limite)]

    col = _achar_coluna(df.columns, col_filtro)
    mantem = df[col].astype(str).str.contains(regex_filtro, case=False, regex=True, na=False)
    df = df[mantem].reset_index(drop=True)

    if ultima_cfg:
        ultima = list(df.columns).index(_achar_coluna(df.columns, ultima_cfg))
        df = df.iloc[:, : ultima + 1]

    # Client: a mesma busca da planilha (conta não encontrada -> sem_cliente: "erro" ou vazio, conforme o banco)
    conta = _achar_coluna(df.columns, "Account")
    if ms:  # MS: os últimos 10 caracteres da Account contra o Account Code do De-Para
        busca = lambda v: depara["codigo"].get(_chave_conta(str(v).strip()[-CHARS_CONTA_MS:]), sem_cliente)
    else:
        busca = lambda v: depara["numero"].get(_chave_conta(v), sem_cliente)
    df.insert(0, "Client", df[conta].map(busca))
    return df


def tratar_banco(banco: str, arquivos: list[Path], depara: dict[str, dict[str, str]]) -> pd.DataFrame:
    """Trata os arquivos do banco (um, ou os dois do MS) e junta o resultado numa tabela só."""
    cfg = BANCOS[banco]
    # Os arquivos trazem liquidações futuras (ex.: D+2), então a janela (D-2) parte da data do arquivo, não da maior
    # Settlement Date da planilha. Nome sem data: usa a data mais recente da própria planilha.
    achou = re.search(r"(\d{8})(?:_\w+)?\.xlsx$", arquivos[0].name)
    mais_recente = pd.Timestamp(datetime.strptime(achou.group(1), "%Y%m%d")) if achou else None
    tabelas = [tratar(a, depara, cfg, mais_recente) for a in arquivos]
    if len(tabelas) == 1:
        return tabelas[0]
    tabelas = [t for t in tabelas if len(t)] or tabelas[:1]  # tabela vazia não entra na junção
    return pd.concat(tabelas, ignore_index=True).dropna(axis=1, how="all") if len(tabelas) > 1 else tabelas[0]


def para_excel(abas: dict[str, pd.DataFrame]) -> bytes:
    """Uma aba por banco (como no modelo)."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl", datetime_format="yyyy-mm-dd") as w:
        for nome, df in abas.items():
            df.to_excel(w, sheet_name=nome, index=False)
            for coluna in w.sheets[nome].iter_cols(min_row=2):
                for celula in coluna:
                    if isinstance(celula.value, (datetime, pd.Timestamp)):
                        celula.number_format = FORMATO_DATA_EXCEL
    return buf.getvalue()


def pasta_do_mes(d: date, banco: str) -> Path:
    """<raiz>/<ano>/<Mês>/<pasta do banco>"""
    return RAIZ_ARQUIVOS / str(d.year) / MESES[d.month - 1] / BANCOS[banco]["pasta"]


def aviso_pasta(d: date, banco: str | None = None) -> str | None:
    """Diz qual parte do caminho <raiz>/<ano>/<Mês>/<banco> não existe (None se está tudo no lugar)."""
    ano, mes = RAIZ_ARQUIVOS / str(d.year), MESES[d.month - 1]
    try:
        if not RAIZ_ARQUIVOS.exists():
            return f"Não consegui acessar a pasta dos arquivos ({RAIZ_ARQUIVOS}). Confira se o Drive G: está disponível."
        if not ano.exists():
            return f"A pasta do ano {d.year} ainda não existe em {RAIZ_ARQUIVOS}."
        if not (ano / mes).exists():
            return f"A pasta {d.year}/{mes} ainda não existe, então não há arquivos para {d:%d/%m/%Y}."
        if banco and not pasta_do_mes(d, banco).exists():
            return f"A pasta do {banco} não existe em {d.year}/{mes}."
    except OSError as e:
        return f"Não consegui verificar a pasta {d.year}/{mes}: {e}"
    return None


def _prefixos(banco: str) -> list[str]:
    cfg = BANCOS[banco]
    return cfg.get("prefixos") or [cfg["prefixo"]]


def arquivos_da_data(d: date, banco: str) -> list[Path]:
    """Caminho do arquivo da data (um por prefixo do banco; pode não existir)."""
    return [pasta_do_mes(d, banco) / f"{pref} - {d:%Y%m%d}.xlsx" for pref in _prefixos(banco)]


def _nomes(arquivos: list[Path]) -> str:
    return ", ".join(a.name for a in arquivos)


def _chave_data(nome: str) -> str:
    achou = re.search(r"\d{8}", nome)  # alguns nomes variam (ex.: "_Off"); ordena pela data
    return (achou.group(0) if achou else "") + nome


@st.cache_data(ttl=60, show_spinner=False)
def listar_arquivos(banco: str) -> dict[str, list[str]]:
    """Arquivos do banco em todos os meses e anos: texto da lista -> caminhos.
    Banco com mais de um arquivo por dia (MS) aparece uma vez por data, com todos os arquivos daquele dia."""
    prefixos = _prefixos(banco)
    achados = {}
    try:
        for p in RAIZ_ARQUIVOS.glob(f"*/*/{BANCOS[banco]['pasta']}/*.xlsx"):
            if p.name.startswith("~$"):
                continue
            if len(prefixos) == 1:
                achados[p.name] = [str(p)]
                continue
            achou = re.fullmatch(r"(%s) - (\d{8})\.xlsx" % "|".join(map(re.escape, prefixos)), p.name)
            if achou:
                achados.setdefault(achou.group(2), []).append(str(p))
    except OSError:
        pass
    if len(prefixos) > 1:  # texto da lista: "20260930 [MS, MS2]"
        achados = {f"{dia} [{', '.join(Path(c).name.split(' - ')[0].removeprefix('Transactions_') for c in sorted(cs))}]": sorted(cs)
                   for dia, cs in achados.items()}
    return dict(sorted(achados.items(), key=lambda kv: _chave_data(kv[0]), reverse=True))


def config_datas(df: pd.DataFrame) -> dict:
    """Mostra as colunas de data como data abreviada (dd/mm/aaaa) nas tabelas da tela."""
    return {c: st.column_config.DateColumn(format="DD/MM/YYYY") for c in df.columns
            if pd.api.types.is_datetime64_any_dtype(df[c])}


def tela_banco(banco: str, data_sel: date) -> None:
    arquivos = arquivos_da_data(data_sel, banco)

    # Opção avançada, oculta por padrão: lista os arquivos de todos os meses e anos (a caixa aceita digitar para filtrar)
    with st.expander("Escolher outro arquivo da pasta"):
        existentes = listar_arquivos(banco)
        escolhido = st.selectbox(
            "Arquivos da pasta (digite a data para filtrar, ex.: 20260930)", list(existentes), index=None,
            placeholder="Usar o arquivo da data selecionada", key=f"arquivo_{banco}",
        )
        if escolhido:
            arquivos = [Path(c) for c in existentes[escolhido]]

    st.caption("Arquivo: " + " | ".join(str(a) for a in arquivos))
    achados = [a for a in arquivos if a.exists()]
    if not achados:
        st.warning(aviso_pasta(data_sel, banco)
                   or "Arquivo não encontrado para essa data. Confira a data ou se a pasta está acessível.")
        return
    if len(achados) < len(arquivos):
        st.info("Não encontrado: " + _nomes([a for a in arquivos if a not in achados]))
    try:
        df = tratar_banco(banco, achados, carregar_depara())
    except (KeyError, ValueError, StopIteration) as e:
        st.error(str(e))
    except OSError as e:
        st.error(f"Não foi possível abrir o arquivo ou o De-Para ({ARQUIVO_MODELO.name}): {e}")
    else:
        st.success(f"{len(df)} linhas e {len(df.columns)} colunas.")
        st.dataframe(df, use_container_width=True, hide_index=True, column_config=config_datas(df))
        st.download_button(
            "Baixar Excel tratado", para_excel({banco: df}),
            file_name=f"Tratado_{achados[0].name}" if len(achados) == 1 else f"Tratado_{banco}_{_chave_data(achados[0].name)[:8]}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key=f"baixar_{banco}",
        )


def cartoes_html(resultados: dict[str, pd.DataFrame], arquivos: dict[str, list[Path]]) -> str:
    """Cartões (padrão do portal) com as linhas de cada banco e quantas ficaram sem cliente no De-Para."""
    itens = []
    for banco, df in resultados.items():
        sem_cliente = int(df["Client"].isin(["", "erro"]).sum())
        if sem_cliente:
            cls, sub = "div", f"{sem_cliente} sem cliente no De-Para"
        elif len(df):
            cls, sub = "ok", "todos com cliente"
        else:
            cls, sub = "vazio", "nenhuma linha na janela"
        itens.append(
            f'<div class="kpi {cls}"><div class="kpi-label">{html.escape(banco)}</div>'
            f'<div class="kpi-valor">{len(df)}</div>'
            f'<div class="kpi-sub">{sub}</div>'
            f'<div class="kpi-sub">{html.escape(_nomes(arquivos[banco]))}</div></div>'
        )
    total = sum(len(df) for df in resultados.values())
    itens.insert(0, f'<div class="kpi"><div class="kpi-label">Total</div><div class="kpi-valor">{total}</div>'
                    f'<div class="kpi-sub">{len(resultados)} banco(s)</div></div>')
    return '<div class="kpis">' + "".join(itens) + "</div>"


# ---------------- E-mail do Daily Report ----------------
# Seções na ordem do modelo (aba E-mail): título -> bancos do app que entram nela. Seção sem linhas não vai no e-mail.
SECOES_EMAIL = [
    ("BTG Cayman", []),  # ainda fora do app
    ("BTG US", ["BTG US"]),
    ("Morgan Stanley", ["MS"]),
    ("Bradesco Miami", ["Bradesco Miami"]),
    ("Bradesco Europa", []),  # ainda fora do app
]
ABA_EMAIL = "E-mail"
ESTILO_TH = ("background:#6A0F2B;color:#fff;font-weight:bold;padding:4px 8px;border:1px solid #ccc;"
             "text-align:left;white-space:nowrap;")
ESTILO_TD = "padding:3px 8px;border:1px solid #ddd;vertical-align:top;"


@st.cache_data(ttl=300, show_spinner=False)
def carregar_destinatarios() -> tuple[str, str]:
    """Para e CC, lidos da aba E-mail do modelo (linha 2, colunas A e B)."""
    wb = load_workbook(ARQUIVO_MODELO, read_only=True, data_only=True)
    try:
        linha = next(wb[ABA_EMAIL].iter_rows(min_row=2, max_row=2, values_only=True))
    finally:
        wb.close()
    return str(linha[0] or "").strip(), str(linha[1] or "").strip()


def _celula_email(v) -> str:
    if v is None or v is pd.NaT or (isinstance(v, float) and pd.isna(v)):
        return ""
    if isinstance(v, (datetime, pd.Timestamp)):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float):
        texto = f"{v:,.4f}".rstrip("0")
        inteiro, _, dec = texto.partition(".")
        return f"{inteiro}.{dec.ljust(2, '0')}"
    return html.escape(str(v)).replace("\n", " ")


def tabela_email_html(df: pd.DataFrame) -> str:
    cab = "".join(f'<th style="{ESTILO_TH}">{html.escape(str(c))}</th>' for c in df.columns)
    linhas = ""
    for linha in df.itertuples(index=False, name=None):
        celulas = ""
        for v in linha:
            numero = isinstance(v, (int, float)) and not isinstance(v, bool)
            celulas += f'<td style="{ESTILO_TD}{"text-align:right;" if numero else ""}">{_celula_email(v)}</td>'
        linhas += f"<tr>{celulas}</tr>"
    return ('<table style="border-collapse:collapse;font-family:Calibri,Arial,sans-serif;font-size:10pt;">'
            f"<tr>{cab}</tr>{linhas}</table>")


def corpo_email_html(resultados: dict[str, pd.DataFrame]) -> tuple[str, list[str]]:
    """Corpo do e-mail e os títulos das seções incluídas (só as que têm linhas)."""
    partes = ["<p>Bom dia!</p>", "<p>Liquidações em andamento:</p>"]
    incluidas = []
    for titulo, bancos in SECOES_EMAIL:
        tabelas = [resultados[b] for b in bancos if b in resultados and len(resultados[b])]
        if not tabelas:
            continue
        df = pd.concat(tabelas, ignore_index=True) if len(tabelas) > 1 else tabelas[0]
        partes.append(f"<p><b>&middot; {html.escape(titulo)}</b></p>{tabela_email_html(df)}<br>")
        incluidas.append(titulo)
    partes += ["<p>Se precisarem de algo, estou à disposição!</p>", "<p>Abs,</p>"]
    return "".join(partes), incluidas


def _enderecos(lista: str) -> list[str]:
    """Endereços de uma lista do tipo "Nome <a@b.com>; Outro <c@d.com>;" (só o e-mail de cada um)."""
    achados = []
    for parte in re.split(r"[;,]", lista):
        m = re.search(r"[\w.+'-]+@[\w.-]+\.\w+", parte)
        if m:
            achados.append(m.group(0))
    return achados


def _assinatura_padrao_configurada() -> bool:
    """True se o Outlook já insere assinatura sozinho nas mensagens novas."""
    import winreg
    for versao in ("16.0", "15.0", "14.0"):
        try:
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER, rf"Software\Microsoft\Office\{versao}\Common\MailSettings") as k:
                if str(winreg.QueryValueEx(k, "NewSignature")[0]).strip():
                    return True
        except OSError:
            continue
    return False


def carregar_assinatura() -> tuple[str, list[tuple[str, Path]]] | None:
    """Assinatura salva do Outlook (a mais recente da pasta Signatures) como HTML, mais as imagens a anexar.
    Devolve (html, [(cid, arquivo)]); None se não houver assinatura."""
    pasta = Path(os.environ.get("APPDATA", "")) / "Microsoft" / "Signatures"
    arquivos = sorted(pasta.glob("*.htm"), key=lambda p: p.stat().st_mtime, reverse=True)
    if not arquivos:
        return None
    bruto = arquivos[0].read_bytes()
    charset = re.search(rb"charset=[\"']?([\w-]+)", bruto[:3000], flags=re.IGNORECASE)
    try:
        texto = bruto.decode(charset.group(1).decode() if charset else "cp1252")
    except (LookupError, UnicodeDecodeError):
        texto = bruto.decode("cp1252", errors="replace")

    imagens: list[tuple[str, Path]] = []

    cids: dict[Path, str] = {}

    def trocar(m):  # imagem local da assinatura -> cid (anexada como imagem embutida, uma vez por arquivo)
        caminho = pasta / unquote(m.group(1))
        if caminho.is_file():
            if caminho not in cids:
                cids[caminho] = f"assinatura{len(cids)}@tag"
                imagens.append((cids[caminho], caminho))
            return f'src="cid:{cids[caminho]}"'
        return m.group(0)

    texto = re.sub(r'src="([^"]+)"', trocar, texto)
    estilos = "".join(re.findall(r"<style.*?</style>", texto, flags=re.IGNORECASE | re.DOTALL))
    corpo = re.search(r"<body[^>]*>(.*)</body>", texto, flags=re.IGNORECASE | re.DOTALL)
    return estilos + (corpo.group(1) if corpo else texto), imagens


def abrir_email_outlook(assunto: str, para: str, cc: str, corpo: str):
    """Abre o rascunho no Outlook com os destinatários resolvidos e a assinatura depois do corpo. Não envia.
    Devolve a lista de destinatários que o Outlook não conseguiu resolver."""
    import pythoncom
    import win32com.client as win32

    pythoncom.CoInitialize()
    try:
        mail = win32.Dispatch("outlook.application").CreateItem(0)
        for tipo, lista in ((1, para), (2, cc)):  # 1 = Para, 2 = CC
            for endereco in _enderecos(lista):
                mail.Recipients.Add(endereco).Type = tipo
        mail.Recipients.ResolveAll()
        mail.Subject = assunto
        nao_resolvidos = [r.Name for r in mail.Recipients if not r.Resolved]

        mail.Display()  # carrega a assinatura padrão, se houver uma configurada
        for _ in range(10):  # espera a assinatura carregar antes de inserir o texto
            try:
                if re.search(r"<body", mail.HTMLBody, flags=re.IGNORECASE):
                    break
            except Exception:
                pass
            time.sleep(0.2)

        extra = ""
        if not _assinatura_padrao_configurada():  # Outlook sem assinatura padrão: usa a salva em Signatures
            assinatura = carregar_assinatura()
            if assinatura:
                extra, imagens = assinatura
                for cid, arquivo in imagens:
                    anexo = mail.Attachments.Add(str(arquivo))
                    anexo.PropertyAccessor.SetProperty("http://schemas.microsoft.com/mapi/proptag/0x3712001F", cid)
                    anexo.PropertyAccessor.SetProperty("http://schemas.microsoft.com/mapi/proptag/0x7FFE000B", True)
        atual = mail.HTMLBody
        m = re.search(r"<body[^>]*>", atual, flags=re.IGNORECASE)
        # corpo, depois a assinatura (a do Outlook, se houver, já vem no corpo atual e fica abaixo)
        mail.HTMLBody = (atual[: m.end()] + corpo + extra + atual[m.end():]) if m else corpo + extra + atual
        return nao_resolvidos
    finally:
        pythoncom.CoUninitialize()


def botao_email(resultados: dict[str, pd.DataFrame], data_sel: date) -> None:
    if st.button("Montar e-mail no Outlook", key="montar_email"):
        corpo, incluidas = corpo_email_html(resultados)
        if not incluidas:
            st.warning("Nenhum banco tem linhas para essa data: o e-mail não foi montado.")
            return
        try:
            para, cc = carregar_destinatarios()
            nao_resolvidos = abrir_email_outlook(
                f"{data_sel:%d/%m/%Y} Daily Report | Offshore Transactions", para, cc, corpo)
        except Exception as e:  # Outlook fechado/indisponível ou modelo inacessível
            st.error(f"Não foi possível montar o e-mail. Abra o Outlook e tente de novo. Detalhe: {e}")
        else:
            st.success("E-mail aberto no Outlook (rascunho, não enviado). Seções: " + ", ".join(incluidas) + ".")
            if nao_resolvidos:
                st.warning("O Outlook não reconheceu estes destinatários: " + ", ".join(nao_resolvidos))


def tela_todos(data_sel: date) -> None:
    """Trata todos os bancos da data de uma vez e junta num Excel, uma aba por banco."""
    depara = carregar_depara_seguro()
    if depara is None:
        return

    arquivos = {banco: arquivos_da_data(data_sel, banco) for banco in BANCOS}
    # Opção avançada, oculta por padrão: escolher outro arquivo (de qualquer data) para cada banco
    with st.expander("Escolher outros arquivos da pasta"):
        for col, banco in zip(st.columns(len(BANCOS)), BANCOS):
            with col:
                existentes = listar_arquivos(banco)
                escolhido = st.selectbox(
                    f"{banco} (digite a data para filtrar)", list(existentes), index=None,
                    placeholder="Usar o arquivo da data selecionada", key=f"arquivo_todos_{banco}",
                )
                if escolhido:
                    arquivos[banco] = [Path(c) for c in existentes[escolhido]]

    resultados = {}
    aviso_mes = aviso_pasta(data_sel)  # raiz, ano ou mês inexistente vale para todos os bancos: avisa uma vez só
    pasta_mes_ausente = aviso_mes is not None
    if aviso_mes:
        st.warning(aviso_mes)
    for banco in BANCOS:
        achados = [a for a in arquivos[banco] if a.exists()]
        if not achados:
            motivo = aviso_pasta(data_sel, banco)
            if motivo and not pasta_mes_ausente:
                st.warning(f"{banco}: {motivo}")
            elif not motivo:
                st.warning(f"{banco}: arquivo não encontrado ({_nomes(arquivos[banco])}).")
            continue
        if len(achados) < len(arquivos[banco]):
            st.info(f"{banco}: não encontrado " + _nomes([a for a in arquivos[banco] if a not in achados]))
        arquivos[banco] = achados  # o cartão mostra só o que foi usado
        try:
            resultados[banco] = tratar_banco(banco, achados, depara)
        except (KeyError, ValueError, StopIteration, OSError) as e:
            st.error(f"{banco}: {e}")

    if not resultados:
        return
    st.markdown(cartoes_html(resultados, arquivos), unsafe_allow_html=True)
    for banco, df in resultados.items():
        with st.expander(f"{banco} ({len(df)} linhas)", expanded=bool(len(df))):
            st.dataframe(df, use_container_width=True, hide_index=True, column_config=config_datas(df))
    st.download_button(
        "Baixar Excel", para_excel(resultados), file_name=f"Tratado_Transactions_{data_sel:%Y%m%d}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="baixar_todos",
    )
    botao_email(resultados, data_sel)


def carregar_depara_seguro() -> dict[str, dict[str, str]] | None:
    try:
        return carregar_depara()
    except (OSError, KeyError, ValueError, StopIteration) as e:
        st.error(f"Não foi possível abrir o De-Para ({ARQUIVO_MODELO.name}): {e}")
        return None


topo("Transactions Offshore", "Tratamento dos arquivos de transações por banco", app="bradesco")

data_sel = st.date_input("Data do arquivo", value=date.today(), format="DD/MM/YYYY")
nomes_abas = ["Todos"] + list(BANCOS)
for aba, nome in zip(st.tabs(nomes_abas), nomes_abas):
    with aba:
        if nome == "Todos":
            tela_todos(data_sel)
        else:
            tela_banco(nome, data_sel)
