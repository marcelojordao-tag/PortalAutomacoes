#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Trata a planilha Relatorio_Rentabilidade.xlsx (aba "Ativo") para ficar com os mesmos
valores do PDF "Rentabilidade e Performance" da BRITech.

A aba Ativo tem UMA LINHA POR ATIVO POR DIA. O PDF tem UMA LINHA POR ATIVO no periodo.
As regras para passar de um para o outro sao:

    PL Inicial = ValorFinal do primeiro dia (dia-base)       -> pega 1 dia, NAO soma
    Entradas   = soma de ValorEntrada dos dias apos o dia-base
    Saidas     = soma de ValorSaida dos dias apos o dia-base
    Rendas     = soma de ValorRendas dos dias apos o dia-base
    Rendimento = soma de RendimentoLiquida dos dias apos o dia-base
    PL Final   = ValorFinal do ultimo dia                    -> pega 1 dia, NAO soma
    % Rent     = juros compostos: produto de (1 + retorno do dia) - 1,
                 com retorno do dia = Rendimento do dia / ValorInicial do dia
    % Rent dos totais (Estrategia, SubClasse, Classe) = a mesma conta, feita com as linhas do grupo
    % Rent do Total Geral = rentabilidade da carteira (cota): coluna RentabilidadeCotaBruta
                 da aba "Estrategia" (e o 1,1180% do cabecalho do PDF)

Agrupamentos: linhas com o mesmo nome viram um ativo so (NTN-B tem varios codigos,
Provisao tem varios). Acoes ("Bolsa") usam o CodigoAtivo (ex.: BOVA11). Conta Corrente
e Lancamento ficam de fora (zerados / nao aparecem no PDF).

Uso:
    python tratar_rentabilidade.py                                  (procura Relatorio_Rentabilidade.xlsx)
    python tratar_rentabilidade.py "C:\\caminho\\Relatorio_Rentabilidade.xlsx"
    python tratar_rentabilidade.py arquivo.xlsx -o saida.xlsx --inicio 2026-07-31 --fim 2026-08-31

Dentro de outro codigo (ex.: o app Streamlit):
    from tratar_rentabilidade import tratar
    ativos, layout, avisos = tratar(pd.read_excel("Relatorio_Rentabilidade.xlsx", sheet_name="Ativo"),
                                    rentab_carteira=1.118)   # ou ler_rentabilidade_carteira("Relatorio_Rentabilidade.xlsx")

Precisa de: pandas, numpy, openpyxl.
"""
from __future__ import annotations

import argparse
import re
import sys
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ----------------------------------------------------------------------------
# Configuracao
# ----------------------------------------------------------------------------
ABA_ATIVO = "Ativo"
ARQUIVO_PADRAO = "Relatorio_Rentabilidade.xlsx"
SAIDA_PADRAO = "Rentabilidade_Tratada.xlsx"
# Extrato de conta corrente (PDF convertido em xlsx). Fixo, ao lado deste script, ate puxarmos pela API.
EXTRATO_PADRAO = Path(__file__).resolve().parent / "ReportExtratoContaCorrente.xlsx"
# Extratos por carteira, baixados pelo teste.py (ReportExtratoContaCorrente_<id>.xlsx)
EXTRATOS_PADRAO = Path(__file__).resolve().parent / "extratos"
# Linhas do extrato que so movem dinheiro entre caixa e ativos (nao mudam o PL)
_INTERNO = re.compile(r"^(Resgate|Aplica|Compra|Venda|Pagamento Juros|Juros|Amortiza|Cupom|Vencimento)", re.I)
_CUPOM = re.compile(r"^(Pagamento Juros|Juros|Amortiza|Cupom)", re.I)
EXCLUIR_ESTRATEGIAS = {"Lançamento"}  # movimentos de caixa que nao aparecem no PDF

COLUNAS_NECESSARIAS = [
    "Data", "TipoAtivo", "CodigoAtivo", "Mnemonico", "Estrategia", "ClasseAtivo", "SubClasseAtivo",
    "ValorInicial", "ValorEntrada", "ValorSaida", "ValorRendas", "ValorFinal", "RendimentoLiquida",
]
COLUNAS_NUMERICAS = [
    "ValorInicial", "ValorEntrada", "ValorSaida", "ValorRendas", "ValorFinal", "RendimentoLiquida",
]
CHAVE = ["Classe", "SubClasse", "Estratégia", "Ativo"]
VALORES = ["PL Inicial", "Entradas", "Saídas", "Rendas", "PL Final", "Rendimento"]
COLUNAS_SAIDA = ["Classe", "SubClasse", "Estratégia", "Ativo"] + VALORES + ["% Rent"]

# Ordem de exibicao parecida com a do PDF (o que nao estiver aqui vai para o fim, em ordem alfabetica)
ORDEM_CLASSE = ["RENDA FIXA", "RENDA VARIÁVEL", "ESTRUTURADOS", "CAIXA"]
ORDEM_SUBCLASSE = {
    "RENDA FIXA": ["CAIXA", "RENDA FIXA PÓS", "RENDA FIXA CDI+", "RENDA FIXA INFLAÇÃO", "RENDA FIXA PÓS PREVIDÊNCIA"],
    "RENDA VARIÁVEL": ["RENDA VARIÁVEL", "RENDA VARIÁVEL PREVIDÊNCIA"],
    "ESTRUTURADOS": ["RENDA VARIÁVEL", "RENDA FIXA PÓS PREVIDÊNCIA", "MULTIMERCADO PREVIDÊNCIA"],
    "CAIXA": ["CAIXA"],
}
ORDEM_ESTRATEGIA = [
    "CAIXA TPF FUNDOS SIMPLES DI",
    "RF PRIVADO FUNDO DE CRÉDITO PRIVADO MID YIELD",
    "RF PRIVADO FUNDO DE CRÉDITO PRIVADO HIGH GRADE",
    "RF PRIVADO FUNDO DE CRÉDITO PRIVADO HIGH YIELD",
    "RF PRIVADO DEBÊNTURES COMUNS CDI +",
    "RF TPF NTNB OU TESOURO INFLAÇÃO",
    "RV ETFs INDEXADOS BOVESPA",
    "RV FUNDOS VALOR",
    "RV FUNDOS LONG BIAS",
    "MM MACRO",
    "MM ARBITRAGEM",
    "CAIXA",
]


# ----------------------------------------------------------------------------
# Tratamento dos dados
# ----------------------------------------------------------------------------
def _posicao(lista: list, valor) -> int:
    return lista.index(valor) if valor in lista else len(lista)


def _ordenar(res: pd.DataFrame) -> pd.DataFrame:
    res = res.copy()
    res["_c"] = res["Classe"].map(lambda v: _posicao(ORDEM_CLASSE, v))
    res["_s"] = [_posicao(ORDEM_SUBCLASSE.get(c, []), s) for c, s in zip(res["Classe"], res["SubClasse"])]
    res["_e"] = res["Estratégia"].map(lambda v: _posicao(ORDEM_ESTRATEGIA, v))
    res = res.sort_values(["_c", "Classe", "_s", "SubClasse", "_e", "Estratégia", "Ativo"])
    return res.drop(columns=["_c", "_s", "_e"]).reset_index(drop=True)


def _rent_composta(dados: pd.DataFrame, chaves: list) -> pd.Series:
    """Rentabilidade composta por grupo: produto de (1 + rendimento do dia / valor inicial do dia) - 1."""
    dia = dados.groupby(chaves + ["Data"], as_index=False)[["ValorInicial", "RendimentoLiquida"]].sum()
    retorno = np.divide(dia["RendimentoLiquida"].to_numpy(), dia["ValorInicial"].to_numpy(),
                        out=np.zeros(len(dia)), where=dia["ValorInicial"].to_numpy() != 0)
    dia["fator"] = 1.0 + retorno
    return dia.groupby(chaves)["fator"].prod() - 1.0


def montar_layout(ativos: pd.DataFrame, rent: dict) -> pd.DataFrame:
    """Monta a tabela no formato do PDF: ativos seguidos de Total Estrategia / SubClasse / Classe / Geral."""
    linhas = []

    def somar(df: pd.DataFrame) -> dict:
        return {c: round(float(df[c].sum()), 2) for c in VALORES}

    def total(nivel, classe, sub, est, df, pct):
        linhas.append({"Nível": nivel, "Classe": classe, "SubClasse": sub, "Estratégia": est,
                       "Ativo": nivel, **somar(df), "% Rent": pct})

    for classe, d_classe in ativos.groupby("Classe", sort=False):
        for sub, d_sub in d_classe.groupby("SubClasse", sort=False):
            for est, d_est in d_sub.groupby("Estratégia", sort=False):
                for _, r in d_est.iterrows():
                    linhas.append({"Nível": "Ativo", "Classe": classe, "SubClasse": sub, "Estratégia": est,
                                   "Ativo": r["Ativo"], **{c: r[c] for c in VALORES}, "% Rent": r["% Rent"]})
                total("Total Estratégia", classe, sub, est, d_est, rent["estrategia"].get((classe, sub, est), np.nan))
            total("Total SubClasse", classe, sub, "", d_sub, rent["subclasse"].get((classe, sub), np.nan))
        total("Total Classe", classe, "", "", d_classe, rent["classe"].get(classe, np.nan))
    total("Total Geral", "", "", "", ativos, rent["geral"])
    return pd.DataFrame(linhas)


def conferir_carteira(layout: pd.DataFrame) -> dict:
    """Conferencia: rentabilidade da carteira (cota) x rentabilidade simples dos ativos.

    1) simples   = Rendimento total / PL Inicial total
    2) diferenca = rentabilidade da carteira - simples          (negativo: carteira rendeu menos que os ativos)
    3) em reais  = diferenca * PL Inicial total
    Base: sempre o PL Inicial (nao usar PL medio nem PL final). Sem arredondar nos passos.
    """
    g = layout[layout["Nível"] == "Total Geral"].iloc[0]
    rendimento, pl_inicial, carteira = float(g["Rendimento"]), float(g["PL Inicial"]), float(g["% Rent"])
    simples = rendimento / pl_inicial if pl_inicial else float("nan")
    diferenca = carteira - simples
    return {"rendimento": rendimento, "pl_inicial": pl_inicial, "carteira": carteira,
            "simples": simples, "diferenca": diferenca, "diferenca_reais": diferenca * pl_inicial}


def _num_br(v):
    """Numero do extrato: float direto (credito) ou texto brasileiro, com parenteses = negativo (debito)."""
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not re.fullmatch(r"[(-]?[\d.]+,\d+\)?", s):
        return None
    n = float(s.strip("()-").replace(".", "").replace(",", "."))
    return -n if s[0] in "(-" else n


def ler_extrato(caminho) -> pd.DataFrame:
    """Le o Extrato de Conta Corrente (PDF em xlsx). Retorna Movimento, Descricao, Valor (+ entra, - sai) e Tipo.

    Tipo: 'Interno' (resgate/aplicacao/compra: so troca caixa por ativo), 'Cupom' (juros/amortizacao
    recebidos, que sao as Rendas do ativo) ou 'Externo' (taxas, despesas, rebate: mudam o PL).
    """
    raw = pd.read_excel(caminho, header=None)

    def achar(nome):
        for i in range(len(raw)):
            for j, v in raw.iloc[i].items():
                if isinstance(v, str) and v.strip() == nome:
                    return i, j
        raise ValueError(f"Nao achei a coluna '{nome}' no extrato {Path(caminho).name}.")

    linha, c_deb = achar("Débito")
    c_cred, c_mov, c_desc, c_saldo = achar("Crédito")[1], achar("Movimento")[1], achar("Descrição")[1], achar("Saldo")[1]
    regs = []
    for i in range(linha + 1, len(raw)):
        mov, desc = raw.iat[i, c_mov], raw.iat[i, c_desc]
        if pd.isna(mov) or pd.isna(desc) or isinstance(mov, str):
            continue  # o cabecalho se repete a cada pagina do extrato
        deb = next((n for n in (_num_br(raw.iat[i, j]) for j in range(c_deb, c_cred - 1)) if n is not None), 0.0)
        cred = next((n for n in (_num_br(raw.iat[i, j]) for j in range(c_cred - 1, c_saldo)) if n is not None), 0.0)
        desc = re.sub(r"\s+", " ", str(desc)).strip()
        tipo = "Cupom" if _CUPOM.match(desc) else "Interno" if _INTERNO.match(desc) else "Externo"
        regs.append({"Movimento": pd.Timestamp(mov), "Descricao": desc, "Valor": round(cred + deb, 2), "Tipo": tipo})
    return pd.DataFrame(regs, columns=["Movimento", "Descricao", "Valor", "Tipo"])


def montar_ponte(df_ativo: pd.DataFrame, ativos: pd.DataFrame, extrato: pd.DataFrame) -> dict:
    """Explica a diferenca entre o PL e o rendimento dos ativos:

        PL Final = PL Inicial + Rendimento dos ativos
                 + caixa externo do extrato (taxas, despesas, rebate)
                 + variacao das provisoes (aba Ativo, linhas 'Provisao')
                 - (Rendas da planilha - cupom recebido em caixa)
    O residuo (PL Final da planilha - PL Final calculado) deve ser zero.
    """
    inicio, fim = ativos.attrs["inicio"], ativos.attrs["fim"]

    df = df_ativo.copy()
    df["Data"] = pd.to_datetime(df["Data"])
    df["ValorFinal"] = pd.to_numeric(df["ValorFinal"], errors="coerce").fillna(0.0)
    prov = df[df["Mnemonico"].astype(str).str.strip().eq("Provisão") & ~df["Estrategia"].isin(EXCLUIR_ESTRATEGIAS)].copy()
    descr = prov["DescricaoAtivo"] if "DescricaoAtivo" in prov.columns else pd.Series("Provisão", index=prov.index)
    prov["Nome"] = descr.fillna("Provisão").astype(str).str.replace(r"\s*\(.*?\)", "", regex=True).str.strip()
    ini = prov[prov["Data"] == inicio].groupby("Nome")["ValorFinal"].sum()
    fin = prov[prov["Data"] == fim].groupby("Nome")["ValorFinal"].sum()
    provisoes = [{"nome": n, "ini": round(float(ini.get(n, 0.0)), 2), "fim": round(float(fin.get(n, 0.0)), 2)}
                 for n in sorted(set(ini.index) | set(fin.index))]
    for p in provisoes:
        p["delta"] = round(p["fim"] - p["ini"], 2)
    linha_prov = ativos[ativos["Ativo"] == "Provisão"]
    delta_total = float((linha_prov["PL Final"] - linha_prov["PL Inicial"]).sum())
    outras = round(delta_total - sum(p["delta"] for p in provisoes), 2)
    if abs(outras) > 0.005:
        provisoes.append({"nome": "Outras provisões", "ini": None, "fim": None, "delta": outras})

    ext = extrato[(extrato["Movimento"] > inicio) & (extrato["Movimento"] <= fim)]
    externos = [{"data": r.Movimento, "descricao": r.Descricao, "valor": r.Valor} for r in ext[ext["Tipo"] == "Externo"].itertuples()]
    cupom_caixa = round(float(ext.loc[ext["Tipo"] == "Cupom", "Valor"].sum()), 2)
    rendas = round(float(ativos["Rendas"].sum()), 2)

    ponte = {"pl_inicial": round(float(ativos["PL Inicial"].sum()), 2),
             "rendimento": round(float(ativos["Rendimento"].sum()), 2),
             "externos": externos, "provisoes": provisoes,
             "rendas": rendas, "cupom_caixa": cupom_caixa,
             "pl_final": round(float(ativos["PL Final"].sum()), 2)}
    explicado = (sum(e["valor"] for e in externos) + sum(p["delta"] for p in provisoes) - (rendas - cupom_caixa))
    ponte["explicado"] = round(explicado, 2)
    ponte["residuo"] = round(ponte["pl_final"] - (ponte["pl_inicial"] + ponte["rendimento"] + explicado), 2)
    return ponte


def textos_ponte(p: dict) -> list:
    def br(v):
        return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    linhas = ["Ponte da diferença (PL Final − PL Inicial − Rendimento):"]
    linhas += [f"  Extrato {e['data']:%d/%m} {e['descricao']}: {br(e['valor'])}" for e in p["externos"]]
    linhas += [f"  Provisão {q['nome']}: {br(q['delta'])}" for q in p["provisoes"]]
    linhas.append(f"  Rendas da planilha − cupom recebido em caixa: {br(-(p['rendas'] - p['cupom_caixa']))}")
    linhas.append(f"  Total explicado: {br(p['explicado'])} | Resíduo: {br(p['residuo'])}")
    return linhas


def textos_conferencia(c: dict) -> list:
    """Passo a passo da conferencia, em texto (formato brasileiro)."""
    def br(v, casas=2):
        return f"{v:,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return [
        f"Rendimento total = {br(c['rendimento'])} | PL Inicial total = {br(c['pl_inicial'])} | "
        f"Rentabilidade da carteira = {br(c['carteira'] * 100, 4)}%",
        f"1) Rentabilidade simples dos ativos = {br(c['simples'], 9)} ({br(c['simples'] * 100, 4)}%)",
        f"2) Diferenca = {br(c['diferenca'], 9)} ({br(c['diferenca'] * 100, 4)} p.p.)",
        f"3) Diferenca em reais = {br(c['diferenca_reais'])}",
        f"Os ativos renderam R$ {br(abs(c['diferenca_reais']))} {'a mais' if c['diferenca_reais'] < 0 else 'a menos'} "
        f"do que a rentabilidade da carteira indica (normalmente custos e despesas da propria carteira).",
    ]


def tratar(df_ativo: pd.DataFrame, inicio=None, fim=None, rentab_carteira=None):
    """Transforma a aba Ativo (1 linha por ativo por dia) na visao do PDF (1 linha por ativo).

    Retorna (ativos, layout, avisos):
      ativos  - uma linha por ativo, sem subtotais (bom para usar em outras telas)
      layout  - igual ao PDF, com Total Estrategia / SubClasse / Classe / Geral
      avisos  - lista de textos com o que merece atencao

    rentab_carteira: rentabilidade da carteira no periodo, em % (ex.: 1.118), vinda da coluna
    RentabilidadeCotaBruta da aba "Estrategia". Vira o % Rent do Total Geral. Se nao vier,
    o Total Geral usa a soma dos ativos (pode diferir um pouco da cota).
    """
    faltando = [c for c in COLUNAS_NECESSARIAS if c not in df_ativo.columns]
    if faltando:
        raise ValueError("Faltam colunas na aba Ativo: " + ", ".join(faltando))

    df = df_ativo.copy()
    df["Data"] = pd.to_datetime(df["Data"])
    for c in COLUNAS_NUMERICAS:
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0.0)

    inicio = pd.Timestamp(inicio) if inicio else df["Data"].min()
    fim = pd.Timestamp(fim) if fim else df["Data"].max()
    df = df[(df["Data"] >= inicio) & (df["Data"] <= fim)]
    df = df[~df["Estrategia"].isin(EXCLUIR_ESTRATEGIAS)]
    if df.empty:
        raise ValueError("Nao ha linhas no periodo informado.")

    # Nome do ativo: acoes usam o codigo (BOVA11); o resto usa o mnemonico
    bolsa = df["TipoAtivo"].astype(str).str.strip().str.lower().eq("bolsa")
    df["Ativo"] = np.where(bolsa, df["CodigoAtivo"].astype(str), df["Mnemonico"].astype(str))
    df["Ativo"] = df["Ativo"].str.replace(r"\s+", " ", regex=True).str.strip()
    df = df.rename(columns={"ClasseAtivo": "Classe", "SubClasseAtivo": "SubClasse", "Estrategia": "Estratégia"})
    for c in ("Classe", "SubClasse", "Estratégia"):
        df[c] = df[c].fillna("Não definida").astype(str).str.strip()

    # PL Inicial e PL Final: valor de UM dia (nao soma)
    pl_inicial = df[df["Data"] == inicio].groupby(CHAVE)["ValorFinal"].sum().rename("PL Inicial")
    pl_final = df[df["Data"] == fim].groupby(CHAVE)["ValorFinal"].sum().rename("PL Final")

    # Fluxos e rendimento: SOMA dos dias depois do dia-base
    posteriores = df[df["Data"] > inicio]
    somas = posteriores.groupby(CHAVE).agg(
        Entradas=("ValorEntrada", "sum"),
        Saídas=("ValorSaida", "sum"),
        Rendas=("ValorRendas", "sum"),
        Rendimento=("RendimentoLiquida", "sum"),
    )

    # % Rent: juros compostos (retorno do dia = rendimento / valor inicial do dia)
    rentab = _rent_composta(posteriores, CHAVE).rename("% Rent")

    res = pd.concat([pl_inicial, somas, pl_final, rentab], axis=1).fillna(0.0).reset_index()
    res[VALORES] = res[VALORES].round(2)
    res = res[res[VALORES].abs().sum(axis=1) > 0.005]  # fora ativos zerados (ex.: Conta Corrente)
    ativos = _ordenar(res)[CHAVE + VALORES + ["% Rent"]]

    avisos = []
    fecha = (ativos["PL Inicial"] + ativos["Entradas"] - ativos["Saídas"] - ativos["Rendas"]
             + ativos["Rendimento"] - ativos["PL Final"]).round(2)
    for nome, dif in zip(ativos["Ativo"][fecha.abs() > 1.0], fecha[fecha.abs() > 1.0]):
        avisos.append(f"{nome}: PL Inicial + Entradas - Saidas - Rendas + Rendimento difere do PL Final em {dif:,.2f}")
    if df["Data"].max() < fim:
        avisos.append(f"A planilha termina em {df['Data'].max():%d/%m/%Y}, antes da data final pedida ({fim:%d/%m/%Y}).")

    # % Rent dos totais: mesma conta, feita com todas as linhas de cada grupo
    rent_grupos = {
        "estrategia": _rent_composta(posteriores, ["Classe", "SubClasse", "Estratégia"]).to_dict(),
        "subclasse": _rent_composta(posteriores, ["Classe", "SubClasse"]).to_dict(),
        "classe": _rent_composta(posteriores, ["Classe"]).to_dict(),
    }
    soma_ativos = float(_rent_composta(posteriores.assign(_t=1), ["_t"]).iloc[0])
    if rentab_carteira is not None and not pd.isna(rentab_carteira):
        rent_grupos["geral"] = float(rentab_carteira) / 100.0
    else:
        rent_grupos["geral"] = soma_ativos
        avisos.append("Total Geral: % Rent calculado pela soma dos ativos, porque nao achei RentabilidadeCotaBruta "
                      "na aba Estrategia. Pode diferir um pouco do % do PDF (que usa a cota da carteira).")

    ativos.attrs["inicio"], ativos.attrs["fim"] = inicio, fim
    ativos.attrs["rent_soma_ativos"] = soma_ativos
    layout = montar_layout(ativos, rent_grupos)
    ativos.attrs["conferencia"] = conferir_carteira(layout)
    return ativos, layout, avisos


# ----------------------------------------------------------------------------
# Excel de saida
# ----------------------------------------------------------------------------
FONTE = "Arial"
FMT_VALOR = '#,##0.00;[Red](#,##0.00)'
FMT_PCT = '0.00%;[Red](0.00%)'
COR_CAB, COR_ESTRAT, COR_SUB, COR_CLASSE, COR_GERAL = "C9453B", "FBE9E5", "F5D3CB", "E8B5AA", "C9453B"
_borda = Border(*(Side(style="thin", color="E0C9C3"),) * 4)


def _cabecalho(ws, colunas: list, linha: int = 1) -> None:
    for j, nome in enumerate(colunas, 1):
        c = ws.cell(linha, j, nome)
        c.font = Font(name=FONTE, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=COR_CAB)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _borda


def _escrever(ws, tabela: pd.DataFrame, com_niveis: bool) -> None:
    _cabecalho(ws, COLUNAS_SAIDA)
    cores = {"Total Estratégia": COR_ESTRAT, "Total SubClasse": COR_SUB, "Total Classe": COR_CLASSE, "Total Geral": COR_GERAL}
    for i, reg in enumerate(tabela.to_dict("records"), 2):
        nivel = reg.get("Nível", "Ativo") if com_niveis else "Ativo"
        for j, col in enumerate(COLUNAS_SAIDA, 1):
            v = reg[col]
            if isinstance(v, float) and np.isnan(v):
                v = None
            c = ws.cell(i, j, v)
            c.border = _borda
            c.font = Font(name=FONTE, size=10, bold=nivel != "Ativo",
                          color="FFFFFF" if nivel == "Total Geral" else "000000")
            if col in VALORES:
                c.number_format = FMT_VALOR
            elif col == "% Rent":
                c.number_format = FMT_PCT
            if nivel in cores:
                c.fill = PatternFill("solid", fgColor=cores[nivel])
    larguras = [18, 30, 54, 54] + [17] * len(VALORES) + [11]
    for j, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "E2"
    ws.row_dimensions[1].height = 28


def _bloco_ponte(ws, p: dict, cel, linha: int) -> None:
    """Ponte PL Inicial -> PL Final. Subtotais e resultado em formulas; o residuo deve ser 0,00."""
    ws.cell(linha, 1, "Ponte: o que explica a diferença (PL Inicial → PL Final)").font = Font(name=FONTE, size=12, bold=True, color="8F2A22")
    linha += 1
    for col, nome in zip("ABC", ["Item", "R$", "Fonte"]):
        x = cel(f"{col}{linha}", nome, negrito=True, cor="FFFFFF")
        x.fill = PatternFill("solid", fgColor=COR_CAB)
        x.alignment = Alignment(horizontal="center")
    linha += 1

    def item(texto, valor, fonte="", negrito=False, fmt=FMT_VALOR):
        nonlocal linha
        cel(f"A{linha}", texto, negrito=negrito)
        cel(f"B{linha}", valor, fmt, negrito=negrito)
        cel(f"C{linha}", fonte)
        linha += 1
        return linha - 1

    r_pl = item("PL Inicial", p["pl_inicial"], "Planilha (ValorFinal do dia-base)", True)
    r_rend = item("(+) Rendimento dos ativos", p["rendimento"], "Planilha (RendimentoLiquida)")

    item("Caixa externo (extrato de conta corrente)", None, "ReportExtratoContaCorrente.xlsx", True)
    a = linha
    for e in p["externos"]:
        item(f"   {e['data']:%d/%m} {e['descricao']}", e["valor"], "Extrato")
    r_ext = item("   Subtotal caixa externo", f"=SUM(B{a}:B{linha - 1})" if p["externos"] else 0, "", True)

    item("Variação das provisões (valores a liquidar)", None, "Planilha, linhas 'Provisão' da aba Ativo", True)
    a = linha
    for q in p["provisoes"]:
        fonte = "" if q["ini"] is None else f"{q['ini']:,.2f} → {q['fim']:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        item(f"   {q['nome']}", q["delta"], fonte)
    r_prov = item("   Subtotal provisões", f"=SUM(B{a}:B{linha - 1})" if p["provisoes"] else 0, "", True)

    r_rendas = item("(−) Rendas da planilha − cupom recebido em caixa", -(p["rendas"] - p["cupom_caixa"]),
                    f"Rendas {p['rendas']:,.2f} − extrato {p['cupom_caixa']:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."))

    r_calc = item("PL Final calculado", f"=B{r_pl}+B{r_rend}+B{r_ext}+B{r_prov}+B{r_rendas}", "", True)
    r_fim = item("PL Final (planilha / PDF)", p["pl_final"], "Planilha (ValorFinal do último dia)")
    item("Resíduo (deve ser 0,00)", f"=B{r_fim}-B{r_calc}", "Se não for zero, falta algum lançamento", True)
    linha += 1
    r_exp = item("Diferença explicada (caixa + provisões + rendas)", f"=B{r_ext}+B{r_prov}+B{r_rendas}", "= PL Final − PL Inicial − Rendimento", True)
    r_cota = item("Diferença pela cota (passo 3)", "=B11", "Rentabilidade da cota × PL Inicial")
    item("Arredondamento da cota", f"=B{r_cota}-B{r_exp}", "Cota vem com 6 casas; centavos de diferença")


def _aba_conferencia(ws, c: dict, ponte: dict = None) -> None:
    """Aba separada: rentabilidade da carteira (cota) x ativos. Resultados em formulas."""
    def cel(ref, valor, fmt=None, negrito=False, cor="000000"):
        x = ws[ref]
        x.value = valor
        x.font = Font(name=FONTE, size=10, bold=negrito, color=cor)
        x.border = _borda
        if fmt:
            x.number_format = fmt
        return x

    ws["A1"].value = "Conferência da rentabilidade: carteira (cota) x ativos"
    ws["A1"].font = Font(name=FONTE, size=12, bold=True, color="8F2A22")
    ws["A2"].value = "Base: sempre o PL Inicial, nas duas pontas. Sem arredondar nos passos intermediários."
    ws["A2"].font = Font(name=FONTE, size=9, italic=True)

    cel("A4", "Rendimento total (Total Geral)")
    cel("B4", c["rendimento"], FMT_VALOR)
    cel("A5", "PL Inicial total (Total Geral)")
    cel("B5", c["pl_inicial"], FMT_VALOR)
    cel("A6", "Rentabilidade da carteira (cota)")
    cel("B6", c["carteira"], "0.0000%")

    for col, nome in zip("ABCD", ["Passo", "Decimal / R$", "Em %", "Conta"]):
        x = cel(f"{col}8", nome, negrito=True, cor="FFFFFF")
        x.fill = PatternFill("solid", fgColor=COR_CAB)
        x.alignment = Alignment(horizontal="center")
    cel("A9", "1) Rentabilidade simples dos ativos")
    cel("B9", "=B4/B5", "0.000000000")
    cel("C9", "=B9", "0.0000%")
    cel("D9", "Rendimento total ÷ PL Inicial total")
    cel("A10", "2) Diferença (carteira − ativos)")
    cel("B10", "=B6-B9", "0.000000000")
    cel("C10", "=B10", "0.0000%")
    cel("D10", "Negativo: a carteira rendeu menos que a soma dos ativos")
    cel("A11", "3) Diferença em reais", negrito=True)
    cel("B11", "=B10*B5", FMT_VALOR, negrito=True)
    cel("C11", None)
    cel("D11", "Diferença (decimal) × PL Inicial total")

    ws["A13"].value = textos_conferencia(c)[-1]
    ws["A13"].font = Font(name=FONTE, size=10)
    if ponte:
        _bloco_ponte(ws, ponte, cel, 15)
    ws.column_dimensions["A"].width = 62
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 58


def salvar_excel(ativos: pd.DataFrame, layout: pd.DataFrame, avisos: list, caminho: Path, ponte: dict = None) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Rentabilidade"
    _escrever(ws, layout, com_niveis=True)

    wa = wb.create_sheet("Ativos")
    _escrever(wa, ativos, com_niveis=False)
    wa.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS_SAIDA))}{len(ativos) + 1}"

    _aba_conferencia(wb.create_sheet("Conferencia"), ativos.attrs["conferencia"], ponte)

    wi = wb.create_sheet("Info")
    linhas = [
        ("Tratamento da aba Ativo (Relatorio_Rentabilidade.xlsx)", True),
        (f"Periodo: {ativos.attrs['inicio']:%d/%m/%Y} a {ativos.attrs['fim']:%d/%m/%Y}", False),
        (f"Gerado em {datetime.now():%d/%m/%Y %H:%M}", False),
        ("", False),
        ("PL Inicial = ValorFinal do dia-base (nao soma). PL Final = ValorFinal do ultimo dia (nao soma).", False),
        ("Entradas, Saidas, Rendas e Rendimento = soma dos dias depois do dia-base.", False),
        ("% Rent = juros compostos: produto de (1 + rendimento do dia / valor inicial do dia) - 1.", False),
        ("Linhas de total: os valores somam os ativos. % Rent de Estrategia, SubClasse e Classe = mesma conta composta do grupo.", False),
        ("% Rent do Total Geral = rentabilidade da carteira (RentabilidadeCotaBruta, aba Estrategia), igual ao cabecalho do PDF.", False),
        ("A ordem dos ativos dentro de cada estrategia e alfabetica (no PDF e outra).", False),
        ("", False),
        ("Avisos" if avisos else "Sem avisos: PL Inicial + Entradas - Saidas - Rendas + Rendimento = PL Final em todos os ativos.", True),
    ] + [(a, False) for a in avisos]
    for i, (texto, negrito) in enumerate(linhas, 1):
        c = wi.cell(i, 1, texto)
        c.font = Font(name=FONTE, size=11 if negrito else 10, bold=negrito, color="8F2A22" if negrito else "000000")
    wi.column_dimensions["A"].width = 120
    wb.save(caminho)


# ----------------------------------------------------------------------------
# Analise por carteira: cota x PL x movimentacao de cotistas
#
# Pergunta que esta secao responde:
#   "Se a carteira so tivesse rentabilidade, o PL Final seria igual ao PL Final real?"
#
#   1) PL Final esperado pela cota = PL Inicial x (1 + cota)
#   2) Diferenca                   = PL Final real - esperado
#   3) Se houve aplicacao/resgate de cotas, o PL muda sem passar pela rentabilidade:
#        efeito = soma(sinal x quantidade de cotas) x cota do ultimo dia
#   4) Residuo = PL Final real - (esperado + efeito)   -> o que sobra sem explicacao
#
# Usada pelo Excel (salvar_excel_multi) e pelo app Streamlit (app.py).
# ----------------------------------------------------------------------------
COL_ID = "id_cliente_consulta"
ABA_HIST_COTA = "HistCota"
ABA_OP_COTISTA = "OpCotista"
TOLERANCIA_REAIS = 1.00            # ate isso a diferenca e so arredondamento
TOLERANCIA_EXPLICADA_PCT = 0.0005  # 0,05% do PL Inicial: sobra aceitavel depois de descontar as cotas movimentadas

STATUS_OK = "Sem divergência"
STATUS_EXPLICADA = "Explicada por cotistas"
STATUS_DIVERGENTE = "Divergente"

COLUNAS_OPERACOES = ["Data", "Tipo", "Quantidade", "Valor bruto", "Valor líquido", "IR", "Sinal"]


def fmt_brl(v, casas: int = 2) -> str:
    """R$ no formato brasileiro, com sinal de menos tipografico."""
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return "—"
    if round(v, casas) == 0:
        v = 0.0
    s = f"{abs(v):,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"−R$ {s}" if v < 0 else f"R$ {s}"


def fmt_pct(v, casas: int = 4) -> str:
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return "—"
    return f"{v * 100:.{casas}f}%".replace(".", ",")


def _sinal_operacao(tipo) -> int:
    """+1 aplicacao (entra cota), -1 resgate/come-cotas (sai cota), 0 desconhecido."""
    t = str(tipo).strip().lower()
    if t.startswith("aplic"):
        return 1
    if t.startswith("resgate") or "come cota" in t or "comecota" in t:
        return -1
    return 0


def resumo_cotistas(ops, hist, inicio, fim) -> dict:
    """Movimentacao de cotas da carteira no periodo.

    ops  - aba OpCotista (API OperacaoCotistaAnalitico_PorCarteiraPeriodo), pode ser vazia
    hist - aba HistCota (API BuscaHistoricoCotaDia): cota, PL e quantidade de cotas por dia
    """
    cot = {"disponivel": False, "n_operacoes": 0, "aplicacoes": 0.0, "resgates": 0.0, "fluxo_valor": 0.0,
           "fluxo_qtd": 0.0, "qtd_ini": None, "qtd_fim": None, "cota_ini": None, "cota_fim": None,
           "pl_cota_ini": None, "pl_cota_fim": None, "qtd_confere": None, "desconhecidas": [],
           "operacoes": pd.DataFrame(columns=COLUNAS_OPERACOES), "historico": None}

    if hist is not None and not hist.empty:
        h = hist.copy()
        h["Data"] = pd.to_datetime(h["Data"])
        h = h.sort_values("Data").reset_index(drop=True)
        cot["historico"] = h
        a, b = h[h["Data"] == inicio], h[h["Data"] == fim]
        if not a.empty and not b.empty:
            cot.update(disponivel=True,
                       qtd_ini=float(a.iloc[0]["QuantidadeFechamento"]), qtd_fim=float(b.iloc[0]["QuantidadeFechamento"]),
                       cota_ini=float(a.iloc[0]["CotaFechamento"]), cota_fim=float(b.iloc[0]["CotaFechamento"]),
                       pl_cota_ini=float(a.iloc[0]["PLFechamento"]), pl_cota_fim=float(b.iloc[0]["PLFechamento"]))

    if ops is not None and not ops.empty:
        o = ops.copy()
        o["DataConversao"] = pd.to_datetime(o["DataConversao"])
        o = o[(o["DataConversao"] > inicio) & (o["DataConversao"] <= fim)].copy()
        for c in ("Quantidade", "ValorBruto", "ValorLiquido", "ValorIR"):
            o[c] = pd.to_numeric(o[c], errors="coerce").fillna(0.0) if c in o.columns else 0.0
        o["Sinal"] = o["TipoOperacao"].map(_sinal_operacao)
        cot["desconhecidas"] = sorted(set(o.loc[o["Sinal"] == 0, "TipoOperacao"].astype(str)))
        cot["aplicacoes"] = float(o.loc[o["Sinal"] > 0, "ValorBruto"].sum())
        cot["resgates"] = float(o.loc[o["Sinal"] < 0, "ValorBruto"].sum())
        cot["fluxo_valor"] = cot["aplicacoes"] - cot["resgates"]
        cot["fluxo_qtd"] = float((o["Sinal"] * o["Quantidade"]).sum())
        cot["n_operacoes"] = int(len(o))
        cot["operacoes"] = (o.sort_values("DataConversao")
                            .rename(columns={"DataConversao": "Data", "TipoOperacao": "Tipo", "ValorBruto": "Valor bruto",
                                             "ValorLiquido": "Valor líquido", "ValorIR": "IR"})[COLUNAS_OPERACOES]
                            .reset_index(drop=True))

    if cot["disponivel"]:
        dif_qtd = (cot["qtd_fim"] - cot["qtd_ini"]) - cot["fluxo_qtd"]
        cot["qtd_confere"] = abs(dif_qtd) <= max(1.0, 1e-6 * abs(cot["qtd_ini"]))
    return cot


def resumo_divergencia(ativos: pd.DataFrame, cot: dict = None) -> dict:
    """Confere cota x ativos e o PL Final que a cota implica x PL Final da planilha (com ou sem cotistas)."""
    c = ativos.attrs["conferencia"]
    pl_final = float(ativos["PL Final"].sum())
    pl_cota = c["pl_inicial"] * (1 + c["carteira"])
    r = {**c, "pl_final": pl_final, "pl_final_cota": pl_cota, "dif_pl": pl_final - pl_cota,
         "efeito_cotas": None, "pl_final_ajustado": None, "residuo_cotistas": None}
    if cot and cot["disponivel"]:
        efeito = cot["fluxo_qtd"] * cot["cota_fim"]
        r.update(efeito_cotas=efeito, pl_final_ajustado=pl_cota + efeito,
                 residuo_cotistas=pl_final - (pl_cota + efeito))
    return r


def classificar(r: dict, cot: dict) -> tuple:
    """(status, mensagem em portugues) para a carteira."""
    dif = r["dif_pl"]
    if abs(dif) <= TOLERANCIA_REAIS:
        return STATUS_OK, ("O PL final fechou com a rentabilidade da cota (diferença de "
                           f"{fmt_brl(dif)}). Não houve movimentação relevante de cotas no período.")
    sentido = "saiu" if dif < 0 else "entrou"
    if r["residuo_cotistas"] is None:
        return STATUS_DIVERGENTE, (f"O PL final difere do esperado pela cota em {fmt_brl(dif)} ({sentido} dinheiro por fora da "
                                   "rentabilidade) e não foi possível buscar o histórico de cotas para explicar.")
    limite = max(TOLERANCIA_REAIS, TOLERANCIA_EXPLICADA_PCT * r["pl_inicial"])
    if cot["n_operacoes"] and abs(r["residuo_cotistas"]) <= limite:
        return STATUS_EXPLICADA, (
            f"O PL final difere do esperado pela cota em {fmt_brl(dif)} porque {sentido} dinheiro por movimentação de cotas: "
            f"{cot['n_operacoes']} operações (aplicações {fmt_brl(cot['aplicacoes'])}, resgates {fmt_brl(cot['resgates'])}). "
            f"Depois de descontá-las, sobra {fmt_brl(r['residuo_cotistas'])} "
            f"({fmt_pct(abs(r['residuo_cotistas']) / r['pl_inicial'])} do PL inicial), efeito de datas e tributos.")
    if not cot["n_operacoes"]:
        return STATUS_DIVERGENTE, (f"O PL final difere do esperado pela cota em {fmt_brl(dif)} e não há aplicação nem resgate "
                                   "de cotas no período que explique. Precisa de investigação.")
    return STATUS_DIVERGENTE, (f"Mesmo descontando a movimentação de cotas, sobram {fmt_brl(r['residuo_cotistas'])} "
                               f"({fmt_pct(abs(r['residuo_cotistas']) / r['pl_inicial'])} do PL inicial) sem explicação.")


def montar_ponte_cotistas(ativos: pd.DataFrame, cot: dict) -> dict:
    """Ponte para carteiras com movimentacao de cotas: PL Inicial + Rendimento + cotistas + outros = PL Final."""
    pl_i, rend = float(ativos["PL Inicial"].sum()), float(ativos["Rendimento"].sum())
    pl_f = float(ativos["PL Final"].sum())
    return {"pl_inicial": pl_i, "rendimento": rend, "pl_final": pl_f, "fluxo_valor": cot["fluxo_valor"],
            "operacoes": cot["operacoes"], "outros": pl_f - (pl_i + rend + cot["fluxo_valor"])}


def explicar_diferenca(dif: float) -> str:
    """Texto do passo 'Diferenca' (PL real - PL esperado pela cota), conforme o valor."""
    if abs(dif) <= TOLERANCIA_REAIS:
        return (f"Praticamente zero ({fmt_brl(dif)} é arredondamento): o PL real é igual ao que a rentabilidade explica. "
                "Não saiu nem entrou dinheiro por fora da cota.")
    if dif < 0:
        return (f"O PL real ficou {fmt_brl(abs(dif))} ABAIXO do que a rentabilidade explica. Ou seja, saiu esse valor da carteira "
                "sem passar pela cota (por exemplo, resgate de cotas). Os próximos passos verificam se foi isso.")
    return (f"O PL real ficou {fmt_brl(dif)} ACIMA do que a rentabilidade explica. Ou seja, entrou esse valor na carteira "
            "sem passar pela cota (por exemplo, aplicação de cotas). Os próximos passos verificam se foi isso.")


COLUNAS_CONTRIBUICAO = ["Classe", "SubClasse", "Estratégia", "Ativo", "PL Inicial", "Rendimento", "% Rent", "Contribuição"]


def contribuicao_ativos(ativos: pd.DataFrame) -> pd.DataFrame:
    """Quanto cada ativo contribuiu para a rentabilidade da carteira (do maior para o menor).

    Contribuição = Rendimento do ativo ÷ PL Inicial TOTAL da carteira (em fração; 0,0021 = 0,21 p.p.).
    A soma de todas as contribuições é a rentabilidade simples dos ativos (Rendimento total ÷ PL Inicial total).
    Ativos sem rendimento (provisão, caixa) ficam de fora.
    """
    pl_total = float(ativos["PL Inicial"].sum())
    d = ativos.loc[ativos["Rendimento"].abs() > 0.005, COLUNAS_CONTRIBUICAO[:-1]].copy()
    d["Contribuição"] = d["Rendimento"] / pl_total if pl_total else 0.0
    return d.sort_values("Contribuição", ascending=False).reset_index(drop=True)


def montar_passos(r: dict, cot: dict) -> list:
    """Passo a passo em dados (para o app): grupo, rotulo, valor, tipo, explicacao, destaque."""
    P = []

    def add(grupo, rotulo, valor, tipo, expl, destaque=False):
        P.append({"grupo": grupo, "rotulo": rotulo, "valor": valor, "tipo": tipo, "explicacao": expl, "destaque": destaque})

    g1 = "A cota explica o PL?"
    add(g1, "PL inicial", r["pl_inicial"], "brl", "Valor da carteira no dia-base.")
    add(g1, "Rentabilidade da cota", r["carteira"], "pct", "Quanto a carteira rendeu no período, segundo a cota.")
    add(g1, "PL final esperado pela cota", r["pl_final_cota"], "brl",
        "PL inicial × (1 + cota). É o PL final se não houvesse aplicação nem resgate.")
    add(g1, "PL final real", r["pl_final"], "brl", "Valor da carteira no último dia, direto da planilha.")
    add(g1, "Diferença (PL real − PL esperado)", r["dif_pl"], "brl", explicar_diferenca(r["dif_pl"]), True)

    g2 = "Movimentação de cotistas"
    if not cot["disponivel"]:
        add(g2, "Dados de cotistas", None, "txt", "Histórico de cotas não encontrado. Rode o teste.py de novo.")
        return P
    add(g2, "Aplicações de cotas", cot["aplicacoes"], "brl", "Soma do valor bruto das aplicações no período.")
    add(g2, "Resgates de cotas", cot["resgates"], "brl", "Soma do valor bruto dos resgates no período.")
    add(g2, "Movimento líquido", cot["fluxo_valor"], "brl", "Aplicações − resgates.")
    add(g2, "Cotas no início", cot["qtd_ini"], "qtd", "Quantidade de cotas no dia-base (histórico de cotas).")
    add(g2, "Cotas no fim", cot["qtd_fim"], "qtd", "Quantidade de cotas no último dia.")
    add(g2, "Variação de cotas", cot["qtd_fim"] - cot["qtd_ini"], "qtd", "Quantas cotas entraram (+) ou saíram (−): cotas no fim − cotas no início.")
    add(g2, "Soma das operações em cotas", cot["fluxo_qtd"], "qtd",
        "Deve ser igual à variação da quantidade de cotas: confirma que a variação veio das operações."
        + ("" if cot["qtd_confere"] else " ATENÇÃO: não bate."))
    add(g2, "Cota no último dia", cot["cota_fim"], "cota", "Cota de fechamento do último dia.")
    add(g2, "Efeito dos movimentos no PL", r["efeito_cotas"], "brl",
        "Soma das operações em cotas × cota do último dia.")

    g3 = "Conclusão"
    add(g3, "PL final esperado com cotistas", r["pl_final_ajustado"], "brl", "Esperado pela cota + efeito dos movimentos.")
    add(g3, "Resíduo (o que sobra sem explicação)", r["residuo_cotistas"], "brl",
        "PL final real − esperado com cotistas. Perto de zero = explicado.", True)
    return P


def rent_encadeada(d_ativo: pd.DataFrame, inicio, fim) -> float:
    """Rentabilidade dos ativos encadeada dia a dia: produto de (1 + rendimento do dia / valor inicial do dia) - 1."""
    df = d_ativo[~d_ativo["Estrategia"].isin(EXCLUIR_ESTRATEGIAS)].copy()
    df["Data"] = pd.to_datetime(df["Data"])
    df = df[(df["Data"] > inicio) & (df["Data"] <= fim)]
    for c in ("ValorInicial", "RendimentoLiquida"):
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0.0)
    dia = df.groupby("Data")[["ValorInicial", "RendimentoLiquida"]].sum()
    dia = dia[dia["ValorInicial"] != 0]
    return float((1 + dia["RendimentoLiquida"] / dia["ValorInicial"]).prod() - 1)


def ler_nome_extrato(caminho):
    """Nome do cliente no cabecalho do extrato de conta corrente ('Cliente: ...')."""
    try:
        raw = pd.read_excel(caminho, header=None)
    except Exception:
        return None
    for i in range(min(len(raw), 30)):
        vals = raw.iloc[i].tolist()
        for j, v in enumerate(vals):
            if isinstance(v, str) and v.strip().rstrip(":").lower() == "cliente":
                for w in vals[j + 1:]:
                    if isinstance(w, str) and w.strip():
                        return re.sub(r"\s+", " ", w).strip()
    return None


def analisar_carteira(id_cliente, d_ativo, hist, ops, rentab, inicio=None, fim=None, arq_extrato=None) -> dict:
    """Analise completa de uma carteira. Levanta ValueError se a aba Ativo nao der para tratar."""
    ativos, layout, avisos = tratar(d_ativo, inicio, fim, rentab)
    cot = resumo_cotistas(ops, hist, ativos.attrs["inicio"], ativos.attrs["fim"])
    if not cot["disponivel"]:
        avisos.append("Histórico de cotas (aba HistCota) não encontrado: não dá para explicar divergências por cotistas.")
    elif cot["qtd_confere"] is False:
        avisos.append("A variação da quantidade de cotas não bate com a soma das operações de cotistas.")
    if cot["desconhecidas"]:
        avisos.append("Tipo(s) de operação de cotista não classificados (ignorados): " + ", ".join(cot["desconhecidas"]))

    r = resumo_divergencia(ativos, cot)
    status, mensagem = classificar(r, cot)

    nome, ponte, ponte_cot = None, None, None
    arq = Path(arq_extrato) if arq_extrato else None
    extrato_ok = bool(arq and arq.exists())
    if extrato_ok:
        nome = ler_nome_extrato(arq)
    if cot["n_operacoes"]:
        # Com movimentacao de cotas, o extrato e as provisoes nao fecham sozinhos: usa a ponte de cotistas
        ponte_cot = montar_ponte_cotistas(ativos, cot)
    elif extrato_ok:
        try:
            ponte = montar_ponte(d_ativo, ativos, ler_extrato(arq))
            if abs(ponte["residuo"]) > 0.05:
                avisos.append(f"Ponte: resíduo de {fmt_brl(ponte['residuo'])}. Falta algum lançamento (extrato ou provisão).")
        except ValueError as e:
            avisos.append(f"Ponte não gerada: {e}")
    else:
        avisos.append("Ponte não gerada: extrato de conta corrente não encontrado.")

    # Despesas do extrato (so para a planilha); com cotistas a ponte nao usa o extrato, entao le aqui
    externos = None
    if ponte:
        externos = ponte["externos"]
    elif extrato_ok:
        try:
            ext = ler_extrato(arq)
            ext = ext[(ext["Movimento"] > ativos.attrs["inicio"]) & (ext["Movimento"] <= ativos.attrs["fim"]) & (ext["Tipo"] == "Externo")]
            externos = [{"data": e.Movimento, "descricao": e.Descricao, "valor": e.Valor} for e in ext.itertuples()]
        except ValueError:
            pass
    try:
        encadeada = rent_encadeada(d_ativo, ativos.attrs["inicio"], ativos.attrs["fim"])
    except (KeyError, ValueError):
        encadeada = None

    return {"id": str(id_cliente), "nome": nome or f"Carteira {id_cliente}", "ativos": ativos, "layout": layout,
            "encadeada": encadeada, "externos_extrato": externos,
            "avisos": avisos, "resumo": r, "cot": cot, "status": status, "mensagem": mensagem,
            "passos": montar_passos(r, cot), "ponte": ponte, "ponte_cotistas": ponte_cot,
            "contribuicoes": contribuicao_ativos(ativos)}


def _ler_aba(entrada, aba):
    try:
        return pd.read_excel(entrada, sheet_name=aba)
    except ValueError:
        return None


def _do_id(df, id_cliente):
    if df is None:
        return None
    if COL_ID not in df.columns:
        return df.iloc[0:0]
    return df[df[COL_ID].astype(str) == str(id_cliente)]


def analisar_todas(entrada, inicio=None, fim=None, pasta_extratos=None) -> tuple:
    """Analisa todas as carteiras do Relatorio_Rentabilidade.xlsx. Retorna ({id: analise}, {id: erro})."""
    df = pd.read_excel(entrada, sheet_name=ABA_ATIVO)
    if COL_ID not in df.columns:
        raise ValueError(f"A aba '{ABA_ATIVO}' não tem a coluna '{COL_ID}'. Rode o teste.py para gerar o arquivo.")
    hist, ops = _ler_aba(entrada, ABA_HIST_COTA), _ler_aba(entrada, ABA_OP_COTISTA)
    pasta = Path(pasta_extratos) if pasta_extratos else EXTRATOS_PADRAO
    analises, erros = {}, {}
    for id_cliente, d_id in df.groupby(df[COL_ID].astype(str), sort=False):
        try:
            analises[id_cliente] = analisar_carteira(
                id_cliente, d_id, _do_id(hist, id_cliente), _do_id(ops, id_cliente),
                ler_rentabilidade_carteira(entrada, id_cliente), inicio, fim,
                pasta / f"ReportExtratoContaCorrente_{id_cliente}.xlsx")
        except ValueError as e:
            erros[id_cliente] = str(e)
    return analises, erros


# ----------------------------------------------------------------------------
# Excel com varias carteiras
# ----------------------------------------------------------------------------
FMT_QTD = '#,##0.000000;[Red](#,##0.000000)'
FMT_COTA = '0.00000000'
COR_STATUS = {STATUS_OK: ("E2F0D9", "375623"), STATUS_EXPLICADA: ("FFF2CC", "7F6000"), STATUS_DIVERGENTE: ("F8CBAD", "C00000")}


def _nota(ws, ref, texto, negrito=False, cor="595959", tamanho=9):
    c = ws[ref]
    c.value = texto
    c.font = Font(name=FONTE, size=tamanho, italic=not negrito, bold=negrito, color=cor)
    return c


def _aba_divergencias(ws, analises: dict) -> None:
    """Tabela-resumo (com formulas e a conta de cada coluna) + passo a passo por carteira."""
    cols = ["Carteira", "Nome", "PL Inicial", "Rendimento ativos", "% Ativos (simples)", "% Cota", "Dif. p.p.",
            "Dif. R$ (cota x ativos)", "PL Final planilha", "PL Final pela cota", "Dif. PL Final R$",
            "Aplicações cotistas R$", "Resgates cotistas R$", "Efeito das cotas no PL R$", "PL Final esperado c/ cotistas",
            "Resíduo R$", "Status"]
    contas = ["ID da carteira", "Cabeçalho do extrato de C/C", "Planilha: ValorFinal no dia-base",
              "Planilha: soma do RendimentoLiquida", "= Rendimento ÷ PL Inicial", "Aba Estrategia: RentabilidadeCotaBruta",
              "= % Cota − % Ativos", "= Dif. p.p. × PL Inicial", "Planilha: ValorFinal no último dia",
              "= PL Inicial × (1 + % Cota)", "= PL Final planilha − PL Final pela cota",
              "API: operações de cotista (valor bruto)", "API: operações de cotista (valor bruto)",
              "= Σ(sinal × cotas) × cota do último dia", "= PL Final pela cota + Efeito das cotas",
              "= PL Final planilha − esperado c/ cotistas",
              f"Sem divergência se |Dif. PL Final| ≤ R$ {TOLERANCIA_REAIS:.2f}; Explicada se há operações e |Resíduo| ≤ "
              f"{TOLERANCIA_EXPLICADA_PCT * 100:.2f}% do PL Inicial"]
    _cabecalho(ws, cols)
    for j, t in enumerate(contas, 1):
        c = ws.cell(2, j, t)
        c.font = Font(name=FONTE, size=8, italic=True, color="595959")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _borda
    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 60

    n = len(analises)
    for i, (id_cliente, a) in enumerate(analises.items(), 3):
        r, cot = a["resumo"], a["cot"]
        tem_cot = cot["disponivel"]
        vals = [id_cliente, a["nome"], r["pl_inicial"], r["rendimento"], f"=D{i}/C{i}", r["carteira"], f"=F{i}-E{i}",
                f"=G{i}*C{i}", r["pl_final"], f"=C{i}*(1+F{i})", f"=I{i}-J{i}",
                cot["aplicacoes"] if tem_cot else None, cot["resgates"] if tem_cot else None,
                r["efeito_cotas"] if tem_cot else None, f"=J{i}+N{i}" if tem_cot else None,
                f"=I{i}-O{i}" if tem_cot else None,
                (f'=IF(ABS(K{i})<={TOLERANCIA_REAIS},"{STATUS_OK}",'
                 f'IF(AND(L{i}+M{i}>0,ABS(P{i})<=MAX({TOLERANCIA_REAIS},{TOLERANCIA_EXPLICADA_PCT}*C{i})),"{STATUS_EXPLICADA}","{STATUS_DIVERGENTE}"))'
                 if tem_cot else f'=IF(ABS(K{i})<={TOLERANCIA_REAIS},"{STATUS_OK}","{STATUS_DIVERGENTE}")')]
        fmts = [None, None, FMT_VALOR, FMT_VALOR, "0.0000%", "0.0000%", "0.0000%", FMT_VALOR, FMT_VALOR, FMT_VALOR,
                FMT_VALOR, FMT_VALOR, FMT_VALOR, FMT_VALOR, FMT_VALOR, FMT_VALOR, None]
        for j, (v, f) in enumerate(zip(vals, fmts), 1):
            c = ws.cell(i, j, v)
            c.border = _borda
            c.font = Font(name=FONTE, size=10, bold=j == len(vals))
            if f:
                c.number_format = f
    ult = n + 2
    for status, (fundo, fonte) in COR_STATUS.items():
        ws.conditional_formatting.add(f"Q3:Q{ult}", CellIsRule(
            operator="equal", formula=[f'"{status}"'], font=Font(name=FONTE, bold=True, color=fonte),
            fill=PatternFill("solid", bgColor=fundo, fgColor=fundo)))
    ws.freeze_panes = "C3"

    # ---- Passo a passo por carteira ----
    linha = n + 5
    ws.cell(linha, 1, "Passo a passo por carteira").font = Font(name=FONTE, size=12, bold=True, color="8F2A22")
    linha += 1
    _nota(ws, f"A{linha}", "Pergunta: se a carteira só tivesse rentabilidade (sem aporte nem resgate de cotas), "
                           "o PL Final seria igual ao PL Final real?")
    linha += 2

    for id_cliente, a in analises.items():
        r, cot, ativos = a["resumo"], a["cot"], a["ativos"]
        for col in "ABC":
            ws[f"{col}{linha}"].fill = PatternFill("solid", fgColor=COR_CAB)
        ws.cell(linha, 1, f"Carteira {id_cliente} · {a['nome']}").font = Font(name=FONTE, size=10, bold=True, color="FFFFFF")
        linha += 1

        def item(rotulo, valor, fmt, expl, negrito=False):
            nonlocal linha
            ws.cell(linha, 1, rotulo).font = Font(name=FONTE, size=10, bold=negrito)
            c = ws.cell(linha, 2, valor)
            c.number_format = fmt
            c.font = Font(name=FONTE, size=10, bold=negrito)
            _nota(ws, f"C{linha}", expl)
            linha += 1
            return linha - 1

        def sub(texto):
            nonlocal linha
            _nota(ws, f"A{linha}", texto, negrito=True, cor="8F2A22", tamanho=10)
            linha += 1

        sub("A cota explica o PL?")
        l_pl = item("1) PL Inicial (dia-base)", r["pl_inicial"], FMT_VALOR, "Valor da carteira no dia-base.")
        l_cota = item("2) Rentabilidade da cota", r["carteira"], "0.0000%", "Quanto a carteira rendeu no período, segundo a cota.")
        l_esp = item("3) PL Final esperado pela cota", f"=B{l_pl}*(1+B{l_cota})", FMT_VALOR,
                     "= (1) × (1 + (2)). O PL Final se NÃO houvesse aporte nem resgate.")
        l_real = item("4) PL Final real (último dia)", r["pl_final"], FMT_VALOR, "Valor da carteira no último dia, direto da planilha.")
        item("5) Diferença = (4) − (3)", f"=B{l_real}-B{l_esp}", FMT_VALOR, explicar_diferenca(r["dif_pl"]), True)

        if cot["disponivel"]:
            sub("Movimentação de cotistas (API: OperacaoCotistaAnalitico_PorCarteiraPeriodo + BuscaHistoricoCotaDia)")
            l_ap = item("6) Aplicações de cotas (R$)", cot["aplicacoes"], FMT_VALOR, "Soma do valor bruto das aplicações no período.")
            l_rg = item("7) Resgates de cotas (R$)", cot["resgates"], FMT_VALOR, "Soma do valor bruto dos resgates no período.")
            item("8) Movimento líquido = (6) − (7)", f"=B{l_ap}-B{l_rg}", FMT_VALOR, "Negativo: saiu dinheiro dos cotistas.")
            l_qi = item("9) Cotas no início", cot["qtd_ini"], FMT_QTD, "Quantidade de cotas no dia-base (BuscaHistoricoCotaDia).")
            l_qf = item("10) Cotas no fim", cot["qtd_fim"], FMT_QTD, "Quantidade de cotas no último dia.")
            l_dq = item("11) Variação de cotas = (10) − (9)", f"=B{l_qf}-B{l_qi}", FMT_QTD, "Quantas cotas entraram (+) ou saíram (−).")
            l_sq = item("12) Soma das operações em cotas", cot["fluxo_qtd"], FMT_QTD,
                        "Deve ser igual a (11): confirma que a variação de cotas veio das operações listadas." +
                        ("" if cot["qtd_confere"] else " ATENÇÃO: não bate."))
            l_ct = item("13) Cota no último dia", cot["cota_fim"], FMT_COTA, "Cota de fechamento do último dia.")
            l_ef = item("14) Efeito dos movimentos no PL = (12) × (13)", f"=B{l_sq}*B{l_ct}", FMT_VALOR,
                        "Quanto as cotas movimentadas pesam no PL final.")
            l_ae = item("15) PL Final esperado c/ cotistas = (3) + (14)", f"=B{l_esp}+B{l_ef}", FMT_VALOR,
                        "O PL final se a carteira só rendesse E só houvesse as operações de cotistas listadas.")
            item("16) Resíduo = (4) − (15)", f"=B{l_real}-B{l_ae}", FMT_VALOR,
                 "O que sobra sem explicação. Perto de zero: divergência explicada (efeito de datas e tributos).", True)
        else:
            sub("Movimentação de cotistas: sem dados (aba HistCota ausente). Rode o teste.py novamente.")

        fundo, fonte = COR_STATUS[a["status"]]
        c = ws.cell(linha, 1, f"Conclusão: {a['status']}. {a['mensagem']}")
        c.font = Font(name=FONTE, size=10, bold=True, color=fonte)
        c.fill = PatternFill("solid", fgColor=fundo)
        for col in "BC":
            ws[f"{col}{linha}"].fill = PatternFill("solid", fgColor=fundo)
        linha += 1
        entradas, saidas = float(ativos["Entradas"].sum()), float(ativos["Saídas"].sum())
        ws.cell(linha, 1, "Referência: Entradas − Saídas dos ativos").font = Font(name=FONTE, size=10)
        ws.cell(linha, 2, entradas - saidas).number_format = FMT_VALOR
        _nota(ws, f"C{linha}", f"Entradas {fmt_brl(entradas)} − Saídas {fmt_brl(saidas)}. Só referência: movimento entre caixa e "
                               "ativos não mexe no PL.")
        linha += 2

    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 34
    for j in range(3, len(cols) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 20
    ws.column_dimensions["Q"].width = 26


def _aba_op_cotistas(ws, analises: dict) -> None:
    """Todas as operacoes de cotistas das carteiras (aplicacoes e resgates de cotas)."""
    cols = ["Carteira", "Nome", "Data", "Tipo", "Quantidade de cotas", "Valor bruto", "Valor líquido", "IR", "Efeito (+/−)"]
    _cabecalho(ws, cols)
    i = 2
    for id_cliente, a in analises.items():
        for _, o in a["cot"]["operacoes"].iterrows():
            vals = [id_cliente, a["nome"], o["Data"].to_pydatetime(), o["Tipo"], float(o["Quantidade"]),
                    float(o["Valor bruto"]), float(o["Valor líquido"]), float(o["IR"]),
                    "Entra" if o["Sinal"] > 0 else "Sai" if o["Sinal"] < 0 else "Não classificado"]
            for j, v in enumerate(vals, 1):
                c = ws.cell(i, j, v)
                c.border = _borda
                c.font = Font(name=FONTE, size=10)
                if j == 3:
                    c.number_format = "dd/mm/yyyy"
                elif j == 5:
                    c.number_format = FMT_QTD
                elif j in (6, 7, 8):
                    c.number_format = FMT_VALOR
            i += 1
    if i == 2:
        ws.cell(2, 1, "Nenhuma aplicação ou resgate de cotas no período.").font = Font(name=FONTE, size=10, italic=True)
    for j, w in enumerate([12, 50, 12, 24, 22, 18, 18, 12, 18], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    if i > 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{i - 1}"


def degraus_contribuicao(a: dict) -> list:
    """Do somatorio das contribuicoes ate a cota, em degraus: [(descricao, tipo, valor R$, fracao do PL inicial)]."""
    pl_i = float(a["resumo"]["pl_inicial"])
    cota = float(a["resumo"]["carteira"])
    soma_ativos = float(a["contribuicoes"]["Contribuição"].sum())
    enc = a.get("encadeada")
    enc = soma_ativos if enc is None else enc
    linhas = [("Soma das contribuições (rentabilidade simples dos ativos)", "Subtotal", soma_ativos * pl_i, soma_ativos),
              ("Diferença de ponderação (o PL variou no período por resgates/aplicações; R$ = bps × PL inicial, não é caixa)", "Ponderação",
               (enc - soma_ativos) * pl_i, enc - soma_ativos),
              ("Rentabilidade dos ativos encadeada dia a dia", "Subtotal", enc * pl_i, enc)]
    # 'Rendimento s/ <ativo>' no extrato e o provento que ja esta no rendimento do ativo: nao conta de novo
    itens = [(f"Extrato {e['data']:%d/%m} {e['descricao']}", "Caixa / despesas", e["valor"])
             for e in (a.get("externos_extrato") or []) if not re.match(r"Rendimento s/", e["descricao"], re.I)]
    p = a.get("ponte")
    if not p:
        lp = a["ativos"][a["ativos"]["Ativo"] == "Provisão"]
        dprov = float((lp["PL Final"] - lp["PL Inicial"]).sum())
        if abs(dprov) > 0.005:
            itens.append(("Variação das provisões (taxas provisionadas e pagas)", "Provisões", dprov))
    if p:
        itens += [(f"Provisão {q['nome']}", "Provisões", q["delta"]) for q in p["provisoes"] if abs(q["delta"]) > 0.005]
        dif_rendas = -(p["rendas"] - p["cupom_caixa"])
        if abs(dif_rendas) > 0.005:
            itens.append(("Rendas da planilha − cupom recebido em caixa", "Caixa / despesas", dif_rendas))
    extra = [(d, cl, v, v / pl_i if pl_i else 0.0) for d, cl, v in itens]
    linhas += extra
    fecha = cota - enc - sum(x[3] for x in extra)
    linhas.append(("Outros até a cota (tributos, taxas não listadas no extrato, defasagem de datas)", "Outros", fecha * pl_i, fecha))
    linhas.append(("Rentabilidade da cota", "Total", cota * pl_i, cota))
    return linhas


def _aba_contribuicoes(ws, analises: dict) -> None:
    """Ranking de contribuicao de cada ativo, por carteira (maior alta primeiro, maior baixa por ultimo)."""
    cols = ["Carteira", "Posição", "Ativo", "Classe", "Estratégia", "PL Inicial", "Rendimento", "% Rent do ativo",
            "Contribuição (bps)", "Contribuição (%)"]
    _cabecalho(ws, cols)
    i = 2
    for id_cliente, a in analises.items():
        for pos, (_, o) in enumerate(a["contribuicoes"].iterrows(), 1):
            vals = [id_cliente, pos, o["Ativo"], o["Classe"], o["Estratégia"], float(o["PL Inicial"]), float(o["Rendimento"]),
                    float(o["% Rent"]), float(o["Contribuição"]) * 10000, float(o["Contribuição"])]
            fmts = [None, "0", None, None, None, FMT_VALOR, FMT_VALOR, "0.0000%;[Red]-0.0000%", "0.00;[Red]-0.00",
                    "0.0000%;[Red]-0.0000%"]
            for j, (v, f) in enumerate(zip(vals, fmts), 1):
                c = ws.cell(i, j, v)
                c.border = _borda
                c.font = Font(name=FONTE, size=10)
                if f:
                    c.number_format = f
            i += 1
        # Secao so da planilha: do somatorio dos ativos ate a cota, em degraus
        linhas = degraus_contribuicao(a)
        for d, cl, v, ctb in linhas:
            vals = [id_cliente, "—", d, cl, "Caixa, despesas e outros", None, v, None, ctb * 10000, ctb]
            fmts = [None, None, None, None, None, FMT_VALOR, FMT_VALOR, None, "0.00;[Red]-0.00", "0.0000%;[Red]-0.0000%"]
            for j, (val, f) in enumerate(zip(vals, fmts), 1):
                c = ws.cell(i, j, val)
                c.border = _borda
                c.font = Font(name=FONTE, size=10, bold=(cl in ("Total", "Subtotal")), italic=(cl not in ("Total", "Subtotal")))
                if f:
                    c.number_format = f
            i += 1
    _nota(ws, f"A{i + 1}", "Contribuição = Rendimento do ativo ÷ PL Inicial total da carteira, em bps (1 bps = 0,01%). A soma das "
                           "contribuições de uma carteira é a rentabilidade simples dos ativos (Rendimento total ÷ PL Inicial total). "
                           "As linhas em itálico 'Caixa, despesas e outros' (só nesta planilha) mostram o que leva da soma dos ativos até a "
                           "rentabilidade da cota; a última linha de cada carteira é a cota.")
    for j, w in enumerate([12, 9, 48, 22, 44, 18, 18, 16, 20, 18], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(i - 1, 2)}"


def _bloco_ponte_cotistas(ws, p: dict, cel, linha: int) -> None:
    """Ponte com movimentacao de cotas: PL Inicial + Rendimento + cotistas + outros efeitos = PL Final."""
    ws.cell(linha, 1, "Ponte: PL Inicial → PL Final (com movimentação de cotistas)").font = Font(name=FONTE, size=12, bold=True, color="8F2A22")
    linha += 1
    for col, nome in zip("ABC", ["Item", "R$", "Fonte"]):
        x = cel(f"{col}{linha}", nome, negrito=True, cor="FFFFFF")
        x.fill = PatternFill("solid", fgColor=COR_CAB)
        x.alignment = Alignment(horizontal="center")
    linha += 1

    def item(texto, valor, fonte="", negrito=False):
        nonlocal linha
        cel(f"A{linha}", texto, negrito=negrito)
        cel(f"B{linha}", valor, FMT_VALOR, negrito=negrito)
        cel(f"C{linha}", fonte)
        linha += 1
        return linha - 1

    r_pl = item("PL Inicial", p["pl_inicial"], "Planilha (ValorFinal do dia-base)", True)
    r_rend = item("(+) Rendimento dos ativos", p["rendimento"], "Planilha (RendimentoLiquida)")
    item("Movimentação de cotistas (valor bruto)", None, "API OperacaoCotistaAnalitico_PorCarteiraPeriodo", True)
    a = linha
    for _, o in p["operacoes"].iterrows():
        item(f"   {o['Data']:%d/%m} {o['Tipo']}", float(o["Valor bruto"]) * (1 if o["Sinal"] >= 0 else -1), "Operação de cotista")
    r_mov = item("   Subtotal cotistas", f"=SUM(B{a}:B{linha - 1})" if len(p["operacoes"]) else 0, "", True)
    r_fim = item("PL Final (planilha)", p["pl_final"], "Planilha (ValorFinal do último dia)", True)
    item("Outros efeitos (tributos, taxas, provisões e defasagem de datas)",
         f"=B{r_fim}-(B{r_pl}+B{r_rend}+B{r_mov})", "O que sobra depois de rendimento e cotistas", True)


def _aba_conferencia_multi(ws, a: dict) -> None:
    """Aba <id>_Conf: conferencia da rentabilidade + ponte (extrato/provisoes ou cotistas)."""
    ponte, ponte_cot = a["ponte"], a["ponte_cotistas"]
    _aba_conferencia(ws, a["ativos"].attrs["conferencia"], ponte)
    if ponte_cot:
        def cel(ref, valor, fmt=None, negrito=False, cor="000000"):
            x = ws[ref]
            x.value = valor
            x.font = Font(name=FONTE, size=10, bold=negrito, color=cor)
            x.border = _borda
            if fmt:
                x.number_format = fmt
            return x
        _bloco_ponte_cotistas(ws, ponte_cot, cel, 15)


def salvar_excel_multi(analises: dict, caminho: Path) -> None:
    """Excel com: Divergencias (resumo + passo a passo), OpCotistas, uma aba por carteira, <id>_Conf e Info."""
    wb = Workbook()
    wb.remove(wb.active)
    _aba_divergencias(wb.create_sheet("Divergencias"), analises)
    _aba_op_cotistas(wb.create_sheet("OpCotistas"), analises)
    _aba_contribuicoes(wb.create_sheet("Contribuicoes"), analises)
    for id_cliente, a in analises.items():
        _escrever(wb.create_sheet(str(id_cliente)[:31]), a["layout"], com_niveis=True)
        if a["ponte"] or a["ponte_cotistas"]:
            _aba_conferencia_multi(wb.create_sheet(f"{id_cliente}_Conf"[:31]), a)

    wi = wb.create_sheet("Info")
    primeiro = next(iter(analises.values()))["ativos"]
    linhas = [
        ("Tratamento da aba Ativo (Relatorio_Rentabilidade.xlsx), uma aba por carteira", True),
        (f"Periodo: {primeiro.attrs['inicio']:%d/%m/%Y} a {primeiro.attrs['fim']:%d/%m/%Y}", False),
        (f"Gerado em {datetime.now():%d/%m/%Y %H:%M}", False),
        ("", False),
        ("Como ler a aba Divergencias", True),
        ("PL Final pela cota = PL Inicial × (1 + cota): o PL se a carteira só tivesse rentabilidade.", False),
        ("Dif. PL Final = PL Final real − PL Final pela cota. Perto de zero: a cota explica o PL.", False),
        ("Se houve aplicação ou resgate de cotas, o PL muda por fora da rentabilidade. Efeito = Σ(sinal × cotas) × cota do último dia.", False),
        ("Resíduo = PL Final real − (PL Final pela cota + Efeito). Perto de zero: a divergência é explicada pelos cotistas.", False),
        (f"Status: {STATUS_OK} (|Dif.| ≤ R$ {TOLERANCIA_REAIS:.2f}); {STATUS_EXPLICADA} (há operações e |Resíduo| ≤ "
         f"{TOLERANCIA_EXPLICADA_PCT * 100:.2f}% do PL Inicial); {STATUS_DIVERGENTE} (demais casos).", False),
        ("", False),
        ("Fontes (API Britech)", True),
        ("Ativos, rentabilidade e cota: BuscaRentabAtivo_Dia e BuscaRentabEstrategia_PeriodoNaoExplodida.", False),
        ("Cotas e PL por dia: Fundo/BuscaHistoricoCotaDia. Aplicações e resgates: Fundo/OperacaoCotistaAnalitico_PorCarteiraPeriodo.", False),
        ("Extrato de conta corrente: Relatorio/RelatorioExtratoContaCorrente (usado na ponte das carteiras sem movimentação de cotas).", False),
        ("", False),
        ("Avisos por carteira", True),
    ]
    for id_cliente, a in analises.items():
        linhas += [(f"{id_cliente}: {t}", False) for t in a["avisos"]]
        if not a["avisos"]:
            linhas.append((f"{id_cliente}: sem avisos", False))
    for i, (texto, negrito) in enumerate(linhas, 1):
        c = wi.cell(i, 1, texto)
        c.font = Font(name=FONTE, size=11 if negrito else 10, bold=negrito, color="8F2A22" if negrito else "000000")
    wi.column_dimensions["A"].width = 130
    wb.save(caminho)


# ----------------------------------------------------------------------------
# Linha de comando
# ----------------------------------------------------------------------------
def ler_rentabilidade_carteira(entrada, id_cliente=None):
    """Rentabilidade da carteira no periodo (em %), da coluna RentabilidadeCotaBruta da aba Estrategia."""
    try:
        d = pd.read_excel(entrada, sheet_name="Estrategia")
    except ValueError:
        return None
    if "RentabilidadeCotaBruta" not in d.columns:
        return None
    if id_cliente is not None and COL_ID in d.columns:
        d = d[d[COL_ID].astype(str) == str(id_cliente)]
    valores = pd.to_numeric(d["RentabilidadeCotaBruta"], errors="coerce").dropna().unique()
    return float(valores[0]) if len(valores) else None


def _achar_entrada(arg) -> Path:
    if arg:
        return Path(arg)
    for pasta in (Path.cwd(), Path(__file__).resolve().parent):
        if (pasta / ARQUIVO_PADRAO).exists():
            return pasta / ARQUIVO_PADRAO
    raise FileNotFoundError(f"Nao achei {ARQUIVO_PADRAO}. Passe o caminho do arquivo como argumento.")


def main(argv=None) -> int:
    p = argparse.ArgumentParser(description="Trata a planilha de rentabilidade para ficar igual ao PDF da BRITech.")
    p.add_argument("entrada", nargs="?", help=f"planilha de entrada (padrao: {ARQUIVO_PADRAO})")
    p.add_argument("-o", "--saida", help=f"arquivo de saida (padrao: {SAIDA_PADRAO} ao lado da entrada)")
    p.add_argument("--inicio", help="data-base, AAAA-MM-DD (padrao: primeira data da planilha)")
    p.add_argument("--fim", help="data final, AAAA-MM-DD (padrao: ultima data da planilha)")
    p.add_argument("--extratos", help=f"pasta com ReportExtratoContaCorrente_<id>.xlsx (padrao: {EXTRATOS_PADRAO.name}/ ao lado do script)")
    p.add_argument("--extrato", help=f"extrato de conta corrente em xlsx (padrao: {EXTRATO_PADRAO.name} ao lado do script)")
    args = p.parse_args(argv)

    try:
        entrada = _achar_entrada(args.entrada)
        df = pd.read_excel(entrada, sheet_name=ABA_ATIVO)
    except FileNotFoundError as e:
        print(e)
        return 1
    except ValueError as e:
        print(f"Nao consegui ler a aba '{ABA_ATIVO}' de {args.entrada or ARQUIVO_PADRAO}: {e}")
        return 1

    if COL_ID in df.columns:
        analises, erros = analisar_todas(entrada, args.inicio, args.fim, args.extratos)
        for id_cliente, msg in erros.items():
            print(f"Carteira {id_cliente}: {msg}")
        if not analises:
            print("Nenhuma carteira pôde ser tratada.")
            return 1
        saida = Path(args.saida) if args.saida else entrada.with_name(SAIDA_PADRAO)
        salvar_excel_multi(analises, saida)
        for id_cliente, a in analises.items():
            r = a["resumo"]
            geral = a["layout"][a["layout"]["Nível"] == "Total Geral"].iloc[0]
            print(f"{id_cliente} ({a['nome']}): {len(a['ativos'])} ativos | PL Final {fmt_brl(geral['PL Final'])} "
                  f"| % Rent {geral['% Rent'] * 100:.4f}% | [{a['status']}]")
            print(f"  dif cota x ativos {fmt_brl(r['diferenca_reais'])} | dif PL Final {fmt_brl(r['dif_pl'])}"
                  + (f" | efeito cotistas {fmt_brl(r['efeito_cotas'])} | resíduo {fmt_brl(r['residuo_cotistas'])}"
                     if r["residuo_cotistas"] is not None else ""))
            if a["ponte"]:
                print(f"  ponte: explicado {fmt_brl(a['ponte']['explicado'])} | resíduo {fmt_brl(a['ponte']['residuo'])}")
            for t in a["avisos"]:
                print("  ATENCAO:", t)
        print(f"Arquivo gerado: {saida}")
        return 0

    try:
        ativos, layout, avisos = tratar(df, args.inicio, args.fim, ler_rentabilidade_carteira(entrada))
    except ValueError as e:
        print(f"Erro: {e}")
        return 1

    ponte = None
    caminho_extrato = Path(args.extrato) if args.extrato else EXTRATO_PADRAO
    if caminho_extrato.exists():
        try:
            ponte = montar_ponte(df, ativos, ler_extrato(caminho_extrato))
        except ValueError as e:
            avisos.append(f"Ponte da diferenca nao gerada: {e}")
        else:
            if abs(ponte["residuo"]) > 0.05:
                avisos.append(f"Ponte da diferenca: residuo de {ponte['residuo']:,.2f}. Falta algum lancamento (extrato ou provisao).")
    else:
        avisos.append(f"Ponte da diferenca nao gerada: nao achei o extrato {caminho_extrato}.")

    saida = Path(args.saida) if args.saida else entrada.with_name(SAIDA_PADRAO)
    salvar_excel(ativos, layout, avisos, saida, ponte)

    geral = layout[layout["Nível"] == "Total Geral"].iloc[0]
    print(f"Periodo: {ativos.attrs['inicio']:%d/%m/%Y} a {ativos.attrs['fim']:%d/%m/%Y} | {len(ativos)} ativos")
    print(f"Total Geral  PL Inicial {geral['PL Inicial']:,.2f} | PL Final {geral['PL Final']:,.2f} | Rendimento {geral['Rendimento']:,.2f} | % Rent {geral['% Rent'] * 100:.4f}%")
    print("\nConferencia da rentabilidade:")
    for t in textos_conferencia(ativos.attrs["conferencia"]):
        print("  " + t)
    if ponte:
        print()
        for t in textos_ponte(ponte):
            print(t)
    for a in avisos:
        print("ATENCAO:", a)
    print(f"Arquivo gerado: {saida}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
