"""Monitor de cotas: a rentabilidade diária de cada carteira está Acima, Abaixo ou Dentro do normal?

A conta é a do script original (mesmas médias, bandas e limites do dia anterior):

  1. Retorno diário = cota bruta de hoje ÷ cota bruta de ontem − 1.
  2. Para calcular média e desvio, os retornos extremos são deixados de fora (filtro de outliers por MAD).
  3. Para cada janela (252, 126, 63 e 21 dias) calcula-se média e desvio com peso exponencial (EWMA,
     meia-vida de 60, 30, 15 e 5 dias). A banda é média ± n desvios (padrão 1,5).
  4. O retorno de HOJE é comparado com a banda calculada ONTEM (para não olhar o futuro):
       acima do limite superior -> "Acima" | abaixo do inferior -> "Abaixo" | entre os dois -> "Dentro".

Diferenças em relação ao script original (todas configuráveis ou documentadas):
  - Piso para o desvio (padrão 0,5 bps/dia): evita que carteiras de retorno quase constante (desvio ~ 0) fiquem
    com uma banda de largura zero e apareçam como "Abaixo" por diferenças de 0,003 p.p.
  - Cota menor ou igual a zero é tratada como dado inválido (o retorno do dia fica em branco) e gera alerta,
    em vez de produzir retornos como −590% ou infinito.
  - Alertas de qualidade (histórico curto, buracos, cota parada, saltos) e de defasagem (cota mais antiga que as demais).

Este módulo não usa Streamlit: o app (app.py) só chama analisar_monitor() e salvar_excel_monitor().
"""
from __future__ import annotations

import io

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

JANELAS = (252, 126, 63, 21)
HALFLIFE = {252: 60, 126: 30, 63: 15, 21: 5}

ACIMA, ABAIXO, DENTRO = "Acima", "Abaixo", "Dentro"
SEM_STATUS = "Sem status"

PARAMS_PADRAO = {
    "n_desvios": 1.5,   # largura da banda, em desvios-padrão
    "n_mads": 4.0,      # filtro de outliers: quantos MADs a partir da mediana
    "mult_mad": 10.0,   # multiplicador do MAD (o script original usa 10)
    "piso_bps": 0.5,    # piso do desvio-padrão diário, em bps (0 = sem piso, como no script original)
}

_ALIASES = {
    "Data": ["Data", "DATA_COTA", "Data Cota"],
    "Id": ["IdCarteira", "ID_CARTEIRA", "Id Carteira", "id_carteira", "Carteira"],
    "Cota": ["CotaBruta", "COTA_BRUTA", "Cota Bruta"],
}


# ----------------------------------------------------------------------------
# Preparação e cálculo
# ----------------------------------------------------------------------------
def normalizar_cotas(df: pd.DataFrame) -> pd.DataFrame:
    """Devolve Id (texto), Data e Cota, ordenados, sem repetição de (Id, Data). Aceita os nomes de colunas da API e da base."""
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]
    achados = {}
    for padrao, candidatos in _ALIASES.items():
        col = next((c for c in candidatos if c in df.columns), None)
        if col is None:
            raise ValueError(f"Não encontrei a coluna de {padrao} (procurei: {', '.join(candidatos)}). "
                             f"Colunas do arquivo: {', '.join(map(str, df.columns))}")
        achados[padrao] = col
    base = pd.DataFrame({
        "Id": df[achados["Id"]].astype(str).str.strip().str.replace(r"\.0$", "", regex=True),
        "Data": pd.to_datetime(df[achados["Data"]], errors="coerce").dt.normalize(),
        "Cota": pd.to_numeric(df[achados["Cota"]], errors="coerce"),
    }).dropna(subset=["Data"])
    base = base[base["Id"].ne("") & base["Id"].ne("nan")]
    base = base.sort_values(["Id", "Data"]).drop_duplicates(["Id", "Data"], keep="last").reset_index(drop=True)
    return base


def _limpar_outliers(serie: pd.Series, n_mads: float, mult_mad: float) -> pd.Series:
    """Deixa em branco (só para o cálculo de média e desvio) os retornos muito distantes da mediana."""
    mediana = serie.median()
    mad = (serie - mediana).abs().median() * mult_mad
    if mad == 0 or pd.isna(mad):
        return serie
    limpa = serie.copy()
    limpa[(serie - mediana).abs() > n_mads * mad] = np.nan
    return limpa


def calcular_serie(cotas: pd.DataFrame, params: dict) -> pd.DataFrame:
    """Série diária de UMA carteira (colunas Data e Cota, em ordem) com retornos, bandas e status de cada janela."""
    d = cotas[["Data", "Cota"]].reset_index(drop=True).copy()
    anterior = d["Cota"].shift(1)
    valido = (d["Cota"] > 0) & (anterior > 0)
    d["Retorno"] = np.where(valido, d["Cota"] / anterior - 1.0, np.nan)

    para_estatistica = _limpar_outliers(d["Retorno"], params["n_mads"], params["mult_mad"])
    piso = params["piso_bps"] / 10000.0

    for j in JANELAS:
        ewm = para_estatistica.ewm(halflife=HALFLIFE[j], min_periods=j // 2)
        media, desvio = ewm.mean(), ewm.std()
        if piso > 0:
            desvio = desvio.clip(lower=piso)
        sup, inf = media + params["n_desvios"] * desvio, media - params["n_desvios"] * desvio
        sup_ant, inf_ant = sup.shift(1), inf.shift(1)  # banda de ontem, aplicada ao retorno de hoje
        d[f"Media_{j}d"], d[f"Std_{j}d"] = media, desvio
        d[f"Lim_Sup_{j}d"], d[f"Lim_Inf_{j}d"] = sup, inf
        d[f"Lim_Sup_{j}d_Ant"], d[f"Lim_Inf_{j}d_Ant"] = sup_ant, inf_ant
        condicoes = [d["Retorno"] > sup_ant, d["Retorno"] < inf_ant,
                     d["Retorno"].notna() & sup_ant.notna() & inf_ant.notna()]
        d[f"Status_{j}d"] = np.select(condicoes, [ACIMA, ABAIXO, DENTRO], default=None)
    return d


def avaliar_qualidade(cotas: pd.DataFrame, serie: pd.DataFrame, data_ref: pd.Timestamp, defasagem: int) -> list:
    """Alertas sobre os dados da carteira: lista de (nivel, texto), nivel em 'erro', 'alerta' ou 'info'."""
    alertas = []
    ruins = cotas[cotas["Cota"] <= 0]
    if len(ruins):
        alertas.append(("erro", f"{len(ruins)} cota(s) menor(es) ou igual(is) a zero (a primeira em {ruins['Data'].min():%d/%m/%Y}). "
                                "Os retornos desses dias foram ignorados."))
    ret = serie["Retorno"].dropna()
    saltos = ret[ret.abs() > 0.5]
    if len(saltos):
        alertas.append(("alerta", f"{len(saltos)} retorno(s) diário(s) acima de 50% (o maior: {saltos.abs().max() * 100:.1f}%). "
                                  "Pode ser evento real ou erro de dado."))
    if len(ret) >= 20 and (ret == 0).mean() > 0.05:
        alertas.append(("info", f"Cota parada (retorno zero) em {(ret == 0).mean() * 100:.1f}% dos dias."))
    n = len(ret)
    sem = [j for j in JANELAS if n < j // 2 + 1]
    if sem:
        alertas.append(("info", f"Histórico curto ({n} retornos): sem status nas janelas de {', '.join(map(str, sem))} dias."))
    intervalos = cotas["Data"].diff().dt.days
    if (intervalos > 5).any():
        i = intervalos.idxmax()
        alertas.append(("alerta", f"Intervalo de {int(intervalos.max())} dias sem cota (até {cotas.loc[i, 'Data']:%d/%m/%Y})."))
    if defasagem > 0:
        alertas.append(("alerta", f"A cota mais recente é de {cotas['Data'].max():%d/%m/%Y}, {defasagem} dia(s) útil(eis) "
                                  f"antes da data de referência ({data_ref:%d/%m/%Y})."))
    return alertas


def _resumir(id_: str, serie: pd.DataFrame, janela: int, atual: bool) -> dict:
    status = serie[f"Status_{janela}d"].dropna()
    contagem = status.value_counts()
    ultimo = serie.iloc[-1]
    st_ultimo = ultimo[f"Status_{janela}d"]
    return {
        "Id": id_, "Janela": janela,
        "Status": st_ultimo if st_ultimo in (ACIMA, ABAIXO, DENTRO) else SEM_STATUS,
        "Atual": atual,
        "Retorno": ultimo["Retorno"],
        "Lim_Inf_Hoje": ultimo[f"Lim_Inf_{janela}d_Ant"], "Lim_Sup_Hoje": ultimo[f"Lim_Sup_{janela}d_Ant"],
        "Media": ultimo[f"Media_{janela}d"], "Std": ultimo[f"Std_{janela}d"],
        "Lim_Inf_Proximo": ultimo[f"Lim_Inf_{janela}d"], "Lim_Sup_Proximo": ultimo[f"Lim_Sup_{janela}d"],
        ABAIXO: int(contagem.get(ABAIXO, 0)), DENTRO: int(contagem.get(DENTRO, 0)), ACIMA: int(contagem.get(ACIMA, 0)),
        "Dias_Com_Status": int(len(status)),
    }


def analisar_monitor(df_cotas: pd.DataFrame, params: dict = None, data_ref=None) -> dict:
    """Roda o monitor em todas as carteiras da base.

    df_cotas  - colunas Data, IdCarteira e CotaBruta (nomes da API ou da base original)
    data_ref  - data de referência para medir defasagem (padrão: a cota mais recente entre as carteiras)
    Retorna {'series': {id: DataFrame}, 'carteiras': {id: info}, 'resumo': DataFrame (id x janela),
             'data_ref', 'params'}.
    """
    p = {**PARAMS_PADRAO, **(params or {})}
    base = normalizar_cotas(df_cotas)
    if base.empty:
        raise ValueError("A base de cotas está vazia.")
    data_ref = pd.Timestamp(data_ref).normalize() if data_ref is not None else base["Data"].max()

    series, carteiras, resumos = {}, {}, []
    for id_, g in base.groupby("Id", sort=False):
        g = g[["Data", "Cota"]].reset_index(drop=True)
        serie = calcular_serie(g, p)
        ultima = g["Data"].iloc[-1]
        defasagem = int(np.busday_count(ultima.date(), data_ref.date())) if ultima < data_ref else 0
        series[id_] = serie
        carteiras[id_] = {"primeira_data": g["Data"].iloc[0], "ultima_data": ultima, "n_linhas": len(g),
                          "defasagem": defasagem, "qualidade": avaliar_qualidade(g, serie, data_ref, defasagem)}
        resumos += [_resumir(id_, serie, j, defasagem == 0) for j in JANELAS]
    return {"series": series, "carteiras": carteiras, "resumo": pd.DataFrame(resumos), "data_ref": data_ref, "params": p}


def resumo_da_janela(res: dict, janela: int) -> pd.DataFrame:
    """Uma linha por carteira para a janela escolhida (para os indicadores e a lista da tela)."""
    r = res["resumo"]
    return r[r["Janela"] == janela].set_index("Id")


# ----------------------------------------------------------------------------
# Excel
# ----------------------------------------------------------------------------
_COR_CAB = "6A0F2B"
_borda = Border(*(Side(style="thin", color="E6DCE0"),) * 4)
_FONTE = "Arial"


def _cabecalho(ws, colunas: list) -> None:
    for j, nome in enumerate(colunas, 1):
        c = ws.cell(1, j, nome)
        c.font = Font(name=_FONTE, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=_COR_CAB)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _borda
    ws.row_dimensions[1].height = 30


def _num(v):
    """Número para o Excel (NaN/infinito vira célula vazia)."""
    if v is None or (isinstance(v, float) and (np.isnan(v) or np.isinf(v))):
        return None
    return float(v) if isinstance(v, (float, np.floating)) else v


def salvar_excel_monitor(res: dict, destino, nomes: dict = None, com_graficos: bool = True, max_graficos: int = 160) -> dict:
    """Grava o Excel do monitor (Resumo, Base_Processada e, se couber, Gráficos). destino: caminho ou BytesIO.

    Retorna {'graficos': n} com a quantidade de gráficos criados (0 se foram omitidos).
    """
    nomes = nomes or {}
    wb = Workbook()

    # --- Resumo ---
    ws = wb.active
    ws.title = "Resumo"
    cols = ["Carteira", "Nome", "Janela (dias)", "Status no último dia", "Retorno do último dia", "Limite inferior (hoje)",
            "Limite superior (hoje)", "Média", "Desvio", "Dias Abaixo", "Dias Dentro", "Dias Acima", "Última data",
            "Defasagem (dias úteis)", "Limite inferior (próximo dia)", "Limite superior (próximo dia)", "Alertas"]
    _cabecalho(ws, cols)
    linha = 2
    for _, r in res["resumo"].iterrows():
        info = res["carteiras"][r["Id"]]
        alertas = " | ".join(t for _, t in info["qualidade"])
        valores = [r["Id"], nomes.get(r["Id"], ""), int(r["Janela"]), r["Status"], _num(r["Retorno"]), _num(r["Lim_Inf_Hoje"]),
                   _num(r["Lim_Sup_Hoje"]), _num(r["Media"]), _num(r["Std"]), int(r[ABAIXO]), int(r[DENTRO]), int(r[ACIMA]),
                   info["ultima_data"].to_pydatetime(), info["defasagem"], _num(r["Lim_Inf_Proximo"]), _num(r["Lim_Sup_Proximo"]),
                   alertas]
        for j, v in enumerate(valores, 1):
            c = ws.cell(linha, j, v)
            c.font = Font(name=_FONTE, size=10)
            c.border = _borda
            if j in (5, 6, 7, 8, 9, 15, 16):
                c.number_format = "0.0000%"
            elif j == 13:
                c.number_format = "dd/mm/yyyy"
        linha += 1
    for j, w in enumerate([11, 40, 9, 16, 14, 14, 14, 11, 11, 9, 9, 9, 12, 12, 14, 14, 90], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(linha - 1, 2)}"

    # --- Base processada (uma carteira depois da outra) ---
    wd = wb.create_sheet("Base_Processada")
    base_cols = ["Carteira", "Data", "Cota", "Retorno"]
    for j in JANELAS:
        base_cols += [f"Media_{j}d", f"Std_{j}d", f"Lim_Sup_{j}d_Ant", f"Lim_Inf_{j}d_Ant", f"Status_{j}d"]
    _cabecalho(wd, base_cols)
    posicao = {c: i + 1 for i, c in enumerate(base_cols)}
    linha, faixas = 2, {}
    for id_, serie in res["series"].items():
        inicio = linha
        for _, s in serie.iterrows():
            wd.append([id_, s["Data"].to_pydatetime(), _num(s["Cota"]), _num(s["Retorno"])] +
                      [x for j in JANELAS for x in (_num(s[f"Media_{j}d"]), _num(s[f"Std_{j}d"]), _num(s[f"Lim_Sup_{j}d_Ant"]),
                                                    _num(s[f"Lim_Inf_{j}d_Ant"]), s[f"Status_{j}d"])])
            linha += 1
        faixas[id_] = (inicio, linha - 1)
    for row in wd.iter_rows(min_row=2, max_row=linha - 1):
        row[1].number_format = "dd/mm/yyyy"
        row[2].number_format = "0.00000000"
        for c in row[3:]:
            if not isinstance(c.value, str):
                c.number_format = "0.0000%"
    for j, w in enumerate([11, 12, 14, 12] + [13, 13, 15, 15, 12] * len(JANELAS), 1):
        wd.column_dimensions[get_column_letter(j)].width = w
    wd.freeze_panes = "C2"

    # --- Gráficos ---
    total = len(res["series"]) * len(JANELAS)
    criados = 0
    if com_graficos and 0 < total <= max_graficos:
        wg = wb.create_sheet("Gráficos")
        topo_linha = 2
        for id_, serie in res["series"].items():
            inicio, _ = faixas[id_]
            for j in JANELAS:
                validas = serie.index[serie[f"Lim_Sup_{j}d_Ant"].notna()]
                if len(validas) == 0:
                    continue
                r0, r1 = inicio + int(validas.min()), inicio + int(validas.max())
                grafico = ScatterChart()
                grafico.title = f"Carteira {id_} · {nomes.get(id_, '')} · janela de {j} dias (EWMA)".replace(" ·  ·", " ·")
                grafico.style = 13
                grafico.height, grafico.width = 12, 24
                grafico.x_axis.title, grafico.y_axis.title = "Data", "Retorno diário"
                grafico.x_axis.number_format, grafico.y_axis.number_format = "dd/mm/yy", "0.00%"
                grafico.x_axis.delete = False  # sem isto o Excel recente esconde os eixos dos gráficos do openpyxl
                grafico.y_axis.delete = False
                grafico.x_axis.tickLblPos = "low"
                grafico.legend.position = "b"
                x = Reference(wd, min_col=posicao["Data"], min_row=r0, max_row=r1)
                pontos = Series(Reference(wd, min_col=posicao["Retorno"], min_row=r0, max_row=r1), x, title="Retorno diário")
                pontos.marker.symbol, pontos.marker.size = "circle", 3
                pontos.marker.graphicalProperties.solidFill = "6A0F2B"
                pontos.marker.graphicalProperties.line.solidFill = "6A0F2B"
                pontos.graphicalProperties.line.noFill = True
                grafico.series.append(pontos)
                for nome_col, titulo, cor in ((f"Lim_Sup_{j}d_Ant", "Limite superior", "2E8B57"),
                                              (f"Lim_Inf_{j}d_Ant", "Limite inferior", "C0392B")):
                    linha_serie = Series(Reference(wd, min_col=posicao[nome_col], min_row=r0, max_row=r1), x, title=titulo)
                    linha_serie.marker.symbol = "none"
                    linha_serie.graphicalProperties.line.solidFill = cor
                    linha_serie.graphicalProperties.line.width = 19050
                    grafico.series.append(linha_serie)
                wg.add_chart(grafico, f"B{topo_linha}")
                topo_linha += 26
                criados += 1
    wb.save(destino)
    return {"graficos": criados, "omitidos": bool(com_graficos and total > max_graficos)}


def excel_monitor_bytes(res: dict, nomes: dict = None, com_graficos: bool = True) -> tuple:
    """Mesmo que salvar_excel_monitor, devolvendo (bytes, info) para download."""
    buf = io.BytesIO()
    info = salvar_excel_monitor(res, buf, nomes, com_graficos)
    return buf.getvalue(), info
